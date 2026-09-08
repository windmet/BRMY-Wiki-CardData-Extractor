import json
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from scripts.compare_exports import compare_xlsx
from toolkit.domains.card_update import merge_card_workbook


HEADERS = ["卡牌", "卡牌编号", "卡牌名", "卡牌译名", "卡牌语音①C"]


def make_old_workbook(path, duplicate=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "cards_data"
    sheet.append(HEADERS + ["人工备注"])
    sheet.append(["card_1", 1, "Old Name", "译名一", "语音翻译", "=1+1"])
    sheet["D2"].number_format = "@"
    sheet["D2"].comment = Comment("keep me", "wiki")
    sheet.append(["card_2", 2, "Removed", "译名二", "", "旧卡"])
    if duplicate:
        sheet.append(["card_1_copy", "001", "Duplicate", "", "", ""])
    workbook.save(path)


class CardUpdateTests(unittest.TestCase):
    def test_styles_and_hyperlinks_survive_workbook_transfer_and_row_reordering(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.xlsx"
            make_old_workbook(old_path)
            workbook = load_workbook(old_path)
            sheet = workbook["cards_data"]
            for coordinate in ("D1", "D2"):
                cell = sheet[coordinate]
                cell.font = Font(name="Microsoft YaHei", bold=True, color="123456")
                cell.fill = PatternFill("solid", fgColor="CDEFFF")
                cell.border = Border(bottom=Side(style="double", color="654321"))
                cell.alignment = Alignment(wrap_text=True, horizontal="right")
                cell.protection = Protection(locked=False)
                cell.number_format = '0.000" custom"'
                cell.quotePrefix = True
            sheet["D2"].hyperlink = "https://example.org/card/1"
            workbook.save(old_path)
            workbook.close()

            merge_card_workbook(
                old_path, HEADERS,
                [["card_3", 3, "Added", "", ""], ["card_1", 1, "Name", "", ""]],
                updated_path=root / "updated.xlsx", changes_path=root / "changes.xlsx",
                audit_path=root / "audit.json",
            )
            old = load_workbook(old_path)
            updated = load_workbook(root / "updated.xlsx")
            for source_coordinate, target_coordinate in (("D1", "D1"), ("D2", "D3")):
                source = old["cards_data"][source_coordinate]
                target = updated["cards_data"][target_coordinate]
                for component in ("font", "fill", "border", "alignment", "protection"):
                    self.assertEqual(copy(getattr(source, component)), copy(getattr(target, component)))
                self.assertEqual(source.number_format, target.number_format)
                self.assertEqual(source.quotePrefix, target.quotePrefix)
            preserved = updated["cards_data"]["D3"]
            self.assertEqual("译名一", preserved.value)
            self.assertEqual("keep me", preserved.comment.text)
            self.assertEqual("https://example.org/card/1", preserved.hyperlink.target)
            self.assertEqual("D3", preserved.hyperlink.ref)
            self.assertIsNone(updated["cards_data"]["D2"].hyperlink)
            old.close()
            updated.close()

    def test_preserves_manual_and_custom_columns_and_builds_diff(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.xlsx"
            updated_path = root / "updated.xlsx"
            changes_path = root / "changes.xlsx"
            audit_path = root / "audit.json"
            make_old_workbook(old_path)

            audit = merge_card_workbook(
                old_path,
                HEADERS,
                [
                    ["card_1", "001", "New Name", "", ""],
                    ["card_3", "3", "Added", "", ""],
                ],
                updated_path=updated_path,
                changes_path=changes_path,
                audit_path=audit_path,
            )

            updated = load_workbook(updated_path, data_only=False)["cards_data"]
            changes = load_workbook(changes_path, data_only=False)
            old = load_workbook(old_path, data_only=False)["cards_data"]
            saved_audit = json.loads(audit_path.read_text(encoding="utf-8"))

        headers = [cell.value for cell in updated[1]]
        indexes = {header: index + 1 for index, header in enumerate(headers)}
        self.assertEqual(1, updated.cell(2, indexes["卡牌编号"]).value)
        self.assertEqual("New Name", updated.cell(2, indexes["卡牌名"]).value)
        self.assertEqual("译名一", updated.cell(2, indexes["卡牌译名"]).value)
        self.assertEqual("语音翻译", updated.cell(2, indexes["卡牌语音①C"]).value)
        self.assertEqual("=1+1", updated.cell(2, indexes["人工备注"]).value)
        self.assertEqual("f", updated.cell(2, indexes["人工备注"]).data_type)
        self.assertEqual("@", updated.cell(2, indexes["卡牌译名"]).number_format)
        self.assertEqual("keep me", updated.cell(2, indexes["卡牌译名"]).comment.text)
        self.assertIsNone(updated.cell(3, indexes["卡牌译名"]).value)
        self.assertEqual(["Summary", "Added", "Modified", "Removed"], changes.sheetnames)
        self.assertEqual(3, changes["Added"]["B2"].value)
        self.assertEqual("卡牌名", changes["Modified"]["B2"].value)
        self.assertEqual(2, changes["Removed"]["B2"].value)
        self.assertEqual("Old Name", old["C2"].value)
        self.assertEqual(audit, saved_audit)
        self.assertEqual(1, audit["Counts"]["AddedCards"])
        self.assertEqual(1, audit["Counts"]["ModifiedCards"])
        self.assertEqual(1, audit["Counts"]["RemovedCards"])
        self.assertEqual(1, audit["Counts"]["PreservedFormulaCells"])
        self.assertEqual(["人工备注"], audit["CustomColumns"])

    def test_duplicate_card_ids_fail_closed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.xlsx"
            make_old_workbook(old_path, duplicate=True)

            with self.assertRaisesRegex(ValueError, "重复卡牌编号"):
                merge_card_workbook(
                    old_path,
                    HEADERS,
                    [["card_1", 1, "Name", "", ""]],
                    updated_path=root / "updated.xlsx",
                    changes_path=root / "changes.xlsx",
                    audit_path=root / "audit.json",
                )

            self.assertFalse((root / "updated.xlsx").exists())

    def test_noop_merge_preserves_blank_cell_types(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "cards_data"
            sheet.append(HEADERS)
            sheet.append(["card_1", 1, "Name", "", ""])
            workbook.save(old_path)

            updated_path = root / "updated.xlsx"
            audit = merge_card_workbook(
                old_path,
                HEADERS,
                [["card_1", "1", "Name", "", ""]],
                updated_path=updated_path,
                changes_path=root / "changes.xlsx",
                audit_path=root / "audit.json",
            )

            difference = compare_xlsx(old_path, updated_path)

        self.assertIsNone(difference)
        self.assertEqual(0, audit["Counts"]["ModifiedCards"])


if __name__ == "__main__":
    unittest.main()
