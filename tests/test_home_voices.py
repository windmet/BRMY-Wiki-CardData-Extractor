import unittest
import tempfile
from pathlib import Path

from toolkit.core.tables import TableCatalog
from toolkit.domains.home_voices import (
    _anomaly_sheet,
    _infer_service_year,
    apply_reference_records,
    build_home_voice_catalog,
    export_home_voice_catalog,
    render_audit_markdown,
)

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def catalog_with(tables):
    names = list(tables)
    return TableCatalog([{name: index for index, name in enumerate(names)}] + [tables[name] for name in names])


class HomeVoiceCatalogTests(unittest.TestCase):
    def setUp(self):
        self.tables = catalog_with({
            "mst_character": [
                {"CharacterId": 1, "CharacterNameJpn": "角色一", "IsActive": True},
                {"CharacterId": 2, "CharacterNameJpn": "角色二", "IsActive": True},
            ],
            "mst_home_voice": [
                {
                    "HomeVoiceTypeCode": 1, "HomeVoiceTargetId": speaker,
                    "HomeVoiceNo": 83, "HomeVoiceCategory": 4,
                    "KeyTargetValue": 2, "MotionCharacterId": speaker,
                    "VoiceCueName": "vo_home_2_83", "IsActive": True,
                }
                for speaker in (1, 2)
            ],
            "mst_season": [],
            "mst_character_home_voice_season": [],
            "mst_character_home_voice_limited": [],
        })

    def test_birthday_subject_uses_cue_target_and_masterdata_year(self):
        scanned = [
            {
                "SpeakerCharacterId": speaker, "AcbFile": f"voice_{speaker}_2.acb",
                "AcbBucket": 2, "CueName": "vo_home_2_83", "CueIndex": 4,
                "CueId": 10, "TitleRaw": title, "TextRaw": f"台词{speaker}",
                "TextWiki": f"台词{speaker}", "MetadataMatchStatus": "matched_by_acb_utf",
                "StableRead": True,
            }
            for speaker, title in ((1, "错误标题"), (2, "角色二の誕生日 [2年目]"))
        ]

        result = build_home_voice_catalog(self.tables, scanned, expected_character_ids=(1, 2))

        self.assertEqual(1, len(result["Subjects"]))
        subject = result["Subjects"][0]
        self.assertEqual("birthday:character=2:year=2", subject["SubjectKey"])
        self.assertEqual("角色二的生日全员祝福 [2年目]", subject["SubjectDisplayName"])
        self.assertTrue(subject["Complete"])
        self.assertEqual("matched", subject["MasterdataStatus"])
        self.assertIn("title_conflict", result["Records"][0]["AuditFlags"])
        self.assertEqual(2, result["Records"][0]["ServiceYear"])
        self.assertEqual("masterdata_key_target", result["Records"][0]["ServiceYearSource"])

    def test_service_year_accepts_the_three_known_marker_families(self):
        first = _infer_service_year({
            "HomeVoiceCategory": 6,
            "CueName": "vo_home_63",
            "TitleRaw": "6月の話題 [1年目]",
            "AcbBucket": 1,
        })
        second = _infer_service_year({
            "HomeVoiceCategory": 6,
            "CueName": "vo_home_106",
            "TitleRaw": "2nd Anniv.",
            "ProductDisplayName": "2nd Anniversary",
            "AcbBucket": 2,
        })
        third = _infer_service_year({
            "HomeVoiceCategory": 4,
            "CueName": "vo_home_4_117",
            "KeyTargetValue": 3,
            "TitleRaw": "綾戸の誕生日 [3年目]",
            "AcbBucket": 3,
        })
        one_and_half = _infer_service_year({
            "HomeVoiceCategory": 6,
            "CueName": "vo_home_99",
            "TitleRaw": "1.5th Anniv.",
            "AcbBucket": 2,
        })

        self.assertEqual((1, "acb_title"), first[:2])
        self.assertEqual((2, "masterdata_product_title"), second[:2])
        self.assertEqual((3, "masterdata_key_target"), third[:2])
        self.assertEqual((2, "acb_package"), one_and_half[:2])

    def test_all_three_birthday_rounds_are_retained(self):
        tables = catalog_with({
            "mst_character": [
                {"CharacterId": 1, "CharacterNameJpn": "キャラクター", "IsActive": True},
            ],
            "mst_home_voice": [
                {
                    "HomeVoiceTypeCode": 1,
                    "HomeVoiceTargetId": 1,
                    "HomeVoiceNo": home_voice_no,
                    "HomeVoiceCategory": 4,
                    "KeyTargetValue": year,
                    "MotionCharacterId": 1,
                    "VoiceCueName": cue_name,
                    "IsActive": True,
                }
                for year, home_voice_no, cue_name in (
                    (1, 30, "vo_home_1_30"),
                    (2, 74, "vo_home_1_74"),
                    (3, 117, "vo_home_1_117"),
                )
            ],
            "mst_season": [],
            "mst_character_home_voice_season": [],
            "mst_character_home_voice_limited": [],
            "mst_home_voice_product": [],
        })
        scanned = [
            {
                "SpeakerCharacterId": 1,
                "AcbFile": f"voice_character_{year}.acb",
                "AcbBucket": year,
                "CueName": cue_name,
                "CueIndex": year,
                "CueId": year,
                "TitleRaw": f"キャラクターの誕生日 [{year}年目]",
                "TextRaw": f"line-{year}",
                "TextWiki": f"line-{year}",
                "MetadataMatchStatus": "matched_by_acb_utf",
                "StableRead": True,
                "AlternateAcbFiles": [],
            }
            for year, cue_name in (
                (1, "vo_home_1_30"),
                (2, "vo_home_1_74"),
                (3, "vo_home_1_117"),
            )
        ]

        result = build_home_voice_catalog(tables, scanned, expected_character_ids=(1,))

        self.assertEqual(
            {
                "birthday:character=1:year=1",
                "birthday:character=1:year=2",
                "birthday:character=1:year=3",
            },
            {subject["SubjectKey"] for subject in result["Subjects"]},
        )

    def test_acb_only_subject_is_retained_and_missing_speaker_is_reported(self):
        scanned = [{
            "SpeakerCharacterId": 1, "AcbFile": "voice_one_general.acb",
            "AcbBucket": 0, "CueName": "vo_home_legacy", "CueIndex": 1,
            "CueId": 1, "TitleRaw": "旧语音", "TextRaw": "文本",
            "TextWiki": "文本", "MetadataMatchStatus": "matched_by_acb_utf",
            "StableRead": True,
        }]

        result = build_home_voice_catalog(self.tables, scanned, expected_character_ids=(1, 2))

        subject = result["Subjects"][0]
        self.assertEqual("acb_only:cue=vo_home_legacy", subject["SubjectKey"])
        self.assertEqual([2], subject["MissingCharacterIds"])
        self.assertFalse(subject["Complete"])
        self.assertEqual("acb_only", result["Records"][0]["MasterdataMatchStatus"])
        self.assertIn("masterdata_missing", result["Records"][0]["InfoFlags"])
        self.assertNotIn("masterdata_missing", result["Records"][0]["AuditFlags"])
        anomaly_rows = _anomaly_sheet(result)["rows"]
        self.assertEqual(1, sum(row[1] == "acb_only_subject" for row in anomaly_rows))
        self.assertEqual(0, sum(row[1] == "masterdata_missing" for row in anomaly_rows))

    def test_non_birthday_key_target_does_not_trigger_package_year_warning(self):
        tables = catalog_with({
            "mst_character": [{"CharacterId": 1, "CharacterNameJpn": "角色一"}],
            "mst_home_voice": [{
                "HomeVoiceTypeCode": 1, "HomeVoiceTargetId": 1, "HomeVoiceNo": 1,
                "HomeVoiceCategory": 9, "KeyTargetValue": 10, "MotionCharacterId": 1,
                "VoiceCueName": "vo_home_1",
            }],
            "mst_season": [],
            "mst_character_home_voice_season": [],
            "mst_character_home_voice_limited": [],
        })
        scanned = [{
            "SpeakerCharacterId": 1, "AcbFile": "voice_one_general.acb",
            "AcbBucket": 0, "CueName": "vo_home_1", "CueIndex": 1, "CueId": 1,
            "TitleRaw": "朝", "TextRaw": "文本", "TextWiki": "文本",
            "MetadataMatchStatus": "matched_by_acb_utf", "StableRead": True,
            "AlternateAcbFiles": [],
        }]

        result = build_home_voice_catalog(tables, scanned, expected_character_ids=(1,))

        self.assertNotIn("package_year_mismatch", result["Records"][0]["AuditFlags"])

    def test_reference_acb_repairs_only_a_duplicated_current_cue(self):
        current = [
            {
                "SpeakerCharacterId": 4, "CueName": cue, "TitleRaw": title,
                "TextRaw": "重复文本", "TextWiki": "重复文本", "AcbFile": "current.acb",
            }
            for cue, title in (("vo_home_6_79", "错误标题"), ("vo_home_7_80", "正确标题"))
        ]
        reference = [
            {
                "SpeakerCharacterId": 4, "CueName": "vo_home_6_79",
                "TitleRaw": "修复标题", "TextRaw": "旧包正确文本",
                "TextWiki": "旧包正确文本", "AcbFile": "reference.acb",
            },
            {
                "SpeakerCharacterId": 4, "CueName": "vo_home_7_80",
                "TitleRaw": "正确标题", "TextRaw": "重复文本",
                "TextWiki": "重复文本", "AcbFile": "reference.acb",
            },
        ]

        repaired, count = apply_reference_records(current, reference)

        self.assertEqual(1, count)
        self.assertEqual("旧包正确文本", repaired[0]["TextRaw"])
        self.assertEqual("重复文本", repaired[0]["OriginalTextRaw"])
        self.assertEqual("reference.acb", repaired[0]["ReferenceAcbFile"])
        self.assertEqual("not_needed", repaired[1]["MetadataRepairStatus"])

    @unittest.skipIf(load_workbook is None, "openpyxl not installed")
    def test_export_is_wiki_facing_and_uses_real_excel_line_breaks(self):
        scanned = [
            {
                "SpeakerCharacterId": speaker, "AcbFile": f"voice_{speaker}_2.acb",
                "AcbBucket": 2, "CueName": "vo_home_2_83", "CueIndex": 4,
                "CueId": 10, "TitleRaw": "角色二の誕生日 [2年目]",
                "TextRaw": f"一行{speaker}\\n二行", "TextWiki": f"一行{speaker}<br>二行",
                "MetadataMatchStatus": "matched_by_acb_utf", "StableRead": True,
                "AlternateAcbFiles": [],
            }
            for speaker in (1, 2)
        ]
        catalog = build_home_voice_catalog(self.tables, scanned, expected_character_ids=(1, 2))

        with tempfile.TemporaryDirectory() as directory:
            paths = export_home_voice_catalog(
                catalog, directory, selected_subject="vo_home_2_83"
            )
            workbook = load_workbook(paths["catalog"], read_only=True)
            subject_workbook = load_workbook(paths["subject"], read_only=True)

            self.assertEqual(["Wiki长表", "完整度"], workbook.sheetnames)
            self.assertEqual(3, workbook["Wiki长表"].max_row)
            self.assertEqual("中文翻译", workbook["Wiki长表"]["E1"].value)
            self.assertEqual("一行1\n二行", workbook["Wiki长表"]["D2"].value)
            self.assertNotIn("<br>", workbook["Wiki长表"]["D2"].value)
            self.assertEqual(2, workbook["完整度"].max_row)
            self.assertEqual(["Wiki主体表"], subject_workbook.sheetnames)
            self.assertEqual(3, subject_workbook["Wiki主体表"].max_row)
            self.assertEqual("一行1\n二行", subject_workbook["Wiki主体表"]["D2"].value)
            workbook.close()
            subject_workbook.close()

            catalog["SourceRoot"] = "current"
            catalog["MasterdataPath"] = "master_data.json"
            catalog["ReferenceAcbRoot"] = "reference"
            catalog["Summary"] = {
                "RecordCount": 2, "SubjectCount": 1, "CompleteSubjectCount": 1,
                "AcbOnlySubjectCount": 0, "ReferenceRepairCount": 0,
            }
            report = render_audit_markdown(catalog)
            self.assertIn("## 数据完整度", report)
            self.assertIn("## 待行动异常", report)
            self.assertIn("2/2", report)


if __name__ == "__main__":
    unittest.main()
