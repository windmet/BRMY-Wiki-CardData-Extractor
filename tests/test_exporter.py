import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from toolkit.core.exporter import write_xlsx


class ExcelTypeTests(unittest.TestCase):
    def test_digit_strings_are_preserved_by_default(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "preserve.xlsx"
            write_xlsx([["00123", "123"]], path, headers=["Code", "Label"])
            sheet = load_workbook(path).active

        self.assertEqual("00123", sheet["A2"].value)
        self.assertEqual("123", sheet["B2"].value)
        self.assertEqual("s", sheet["A2"].data_type)
        self.assertEqual("s", sheet["B2"].data_type)

    def test_integer_conversion_requires_explicit_column_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "integer.xlsx"
            write_xlsx(
                [["00123"]],
                path,
                headers=["Id"],
                column_types={0: "integer"},
            )
            cell = load_workbook(path).active["A2"]

        self.assertEqual(123, cell.value)
        self.assertEqual("n", cell.data_type)
        self.assertEqual("0", cell.number_format)

    def test_text_schema_forces_excel_text_format(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.xlsx"
            write_xlsx(
                [[123]],
                path,
                headers=["Code"],
                column_types={0: "text"},
            )
            cell = load_workbook(path).active["A2"]

        self.assertEqual("123", cell.value)
        self.assertEqual("s", cell.data_type)
        self.assertEqual("@", cell.number_format)

    def test_invalid_integer_value_fails_closed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.xlsx"
            with self.assertRaisesRegex(ValueError, "expects integer"):
                write_xlsx(
                    [["12A"]],
                    path,
                    headers=["Id"],
                    column_types={0: "integer"},
                )
            self.assertFalse(path.exists())


if __name__ == "__main__":
    unittest.main()
