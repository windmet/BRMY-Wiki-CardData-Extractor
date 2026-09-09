import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from toolkit.domains.birthday import (
    _infer_cycle,
    _resolve_cycle,
    _select_character_texts,
    _text_format,
    export,
)


class BirthdayCycleTests(unittest.TestCase):
    def setUp(self):
        self.characters = {
            1: {"Month": 5, "Day": 14},
            2: {"Month": 1, "Day": 6},
        }

    def test_latest_natural_year_stays_in_anniversary_cycle(self):
        rows = [
            {"CharacterId": 1, "Year": 2024},
            {"CharacterId": 2, "Year": 2025},
            {"CharacterId": 1, "Year": 2025},
            {"CharacterId": 2, "Year": 2026},
            {"CharacterId": 1, "Year": 2026},
            {"CharacterId": 2, "Year": 2027},
        ]

        cycle = _resolve_cycle(_infer_cycle(self.characters, rows))

        self.assertEqual(2024, cycle["BaseYear"])
        self.assertEqual((5, 14), (cycle["BoundaryMonth"], cycle["BoundaryDay"]))
        self.assertEqual(2026, cycle["StartYear"])
        self.assertEqual(3, cycle["Cycle"])
        self.assertEqual("auto", cycle["Selection"])

    def test_cycle_override_is_derived_from_base_year(self):
        inferred = {
            "BaseYear": 2024,
            "StartYear": 2026,
            "Cycle": 3,
            "BoundaryMonth": 5,
            "BoundaryDay": 14,
        }

        cycle = _resolve_cycle(inferred, target_cycle=2)

        self.assertEqual(2025, cycle["StartYear"])
        self.assertEqual(2, cycle["Cycle"])
        self.assertEqual("override", cycle["Selection"])

    def test_inconsistent_year_and_cycle_are_rejected(self):
        inferred = {
            "BaseYear": 2024,
            "StartYear": 2026,
            "Cycle": 3,
            "BoundaryMonth": 5,
            "BoundaryDay": 14,
        }

        with self.assertRaisesRegex(ValueError, "不一致"):
            _resolve_cycle(inferred, target_year=2026, target_cycle=2)

        with self.assertRaisesRegex(ValueError, "早于第一轮"):
            _resolve_cycle(inferred, target_year=2023)

    def test_three_known_numbering_formats_are_distinct(self):
        def rows(numbers, targets):
            return [
                {
                    "CharacterBirthdayTextNo": number,
                    "KeyTargetValue": target,
                    "Text": f"line-{number}",
                }
                for number, target in zip(numbers, targets)
            ]

        self.assertEqual("zero_target_offset_numbers", _text_format(rows(range(76, 81), [0] * 5)))
        self.assertEqual("zero_target_local_numbers", _text_format(rows(range(1, 6), [0] * 5)))
        self.assertEqual("key_target_sequence", _text_format(rows(range(1, 4), range(1, 4))))

    def test_global_numbers_are_normalized_without_losing_original_ids(self):
        selected = [
            {"CharacterBirthdayTextNo": number, "KeyTargetValue": 0, "Text": f"text-{number}"}
            for number in range(76, 81)
        ]

        texts, metadata = _select_character_texts({16: {2024: selected}}, 16, 2024)

        self.assertEqual([1, 2, 3, 4, 5], list(texts))
        self.assertEqual("text-76", texts[1])
        self.assertEqual(list(range(76, 81)), metadata["OriginalTextNos"])
        self.assertEqual("available", metadata["Status"])

    def test_repeated_text_is_audited_but_never_dropped(self):
        previous = [
            {"CharacterBirthdayTextNo": number, "KeyTargetValue": 0, "Text": f"same-{index}"}
            for index, number in enumerate(range(76, 81), start=1)
        ]
        repeated = [dict(row) for row in previous]

        texts, metadata = _select_character_texts(
            {16: {2024: previous, 2025: repeated}}, 16, 2025
        )

        self.assertEqual(5, len(texts))
        self.assertEqual("available", metadata["Status"])
        self.assertTrue(metadata["MatchesPreviousYear"])

    def test_export_includes_all_five_lines_for_older_cycles(self):
        data = {
            "TargetCycle": 1,
            "CycleStartYear": 2024,
            "CycleBoundary": {"Month": 5, "Day": 14},
            "Characters": {
                "1": {"Name": "Character", "Month": 6, "Day": 1},
            },
            "BirthdayTexts": {
                "1": {str(index): f"text-{index}" for index in range(1, 6)},
            },
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "json_output"
            output.mkdir()
            (output / "birthday_extract.json").write_text(
                json.dumps(data), encoding="utf-8"
            )
            previous = os.getcwd()
            os.chdir(root)
            try:
                export()
            finally:
                os.chdir(previous)

            sheet = load_workbook(root / "xlsx_output" / "birthday_lines.xlsx").active

        self.assertEqual("庆典台词⑤J", sheet["A12"].value)
        self.assertEqual("text-5", sheet["B12"].value)
        self.assertEqual(13, sheet.max_row)
        self.assertEqual("轮次", sheet["A3"].value)
        self.assertEqual(1, sheet["B3"].value)
        self.assertEqual("2024/06/01", sheet["B2"].value)
        self.assertEqual("庆典台词⑤C", sheet["A13"].value)
        self.assertIsNone(sheet["B13"].value)


if __name__ == "__main__":
    unittest.main()
