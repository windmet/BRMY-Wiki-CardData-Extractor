import unittest
import tempfile
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from toolkit.core.rewards import RewardResolver, REWARD_TARGETS
from toolkit.core.tables import TableCatalog
from toolkit.core.output import OutputContext
from toolkit.domains import events


def catalog(**tables):
    return TableCatalog([dict.fromkeys(tables, [])] + list(tables.values()))


class RewardResolverTests(unittest.TestCase):
    def test_each_reward_type_uses_its_own_target_namespace(self):
        tables = {spec[2]: [{spec[3]: 1, spec[4]: f'name-{code}'}]
                  for code, spec in REWARD_TARGETS.items()}
        tables['mst_title'] = [{'TitleId': 1, 'TitleFileName': 'wrong-title-screen'}]
        resolver = RewardResolver(catalog(**tables))
        for code, spec in REWARD_TARGETS.items():
            with self.subTest(code=code):
                result = resolver.resolve({'RewardTypeCode': code, 'RewardTargetId': 1, 'RewardCount': 2})
                self.assertEqual(f'name-{code}', result['RewardName'])
                self.assertEqual(spec[2], result['TargetTable'])
                self.assertEqual('resolved', result['Resolution'])
                self.assertEqual(2, result['RewardCount'])

    def test_unresolved_and_unnamed_targets_keep_identity_without_guessing(self):
        resolver = RewardResolver(catalog(
            mst_item=[{'ItemId': 1, 'ItemName': 'not an ingredient'}],
            mst_honor=[{'HonorId': 1, 'HonorName': '', 'HonorFileName': 'technical', 'IsActive': False}],
            mst_pin=[{'PinId': 1, 'PinName': 'one'}, {'PinId': 1, 'PinName': 'two'}],
        ))
        for code, target, expected in [(101, 1, 'missing_table'), (2, 2, 'missing_target'),
                                       (6, 1, 'name_unavailable'), (7, 1, 'ambiguous_target'),
                                       (99, 0, 'context_required'), (100, 1, 'context_required'),
                                       (102, 0, 'context_required'), (999, 1, 'unknown_type')]:
            with self.subTest(code=code):
                row = {'RewardTypeCode': code, 'RewardTargetId': target, 'RewardCount': 0}
                result = resolver.resolve(row)
                self.assertEqual(expected, result['Resolution'])
                self.assertEqual(row, result['RawReward'])
                self.assertEqual(0, result['RewardCount'])
                self.assertNotIn('technical', result['RewardName'])
                self.assertNotIn('not an ingredient', result['RewardName'])
                if code == 6:
                    self.assertFalse(result['TargetActive'])
        self.assertEqual(8, len(resolver.issues))

    def test_groups_preserve_order_and_multiple_rows_with_separate_namespaces(self):
        resolver = RewardResolver(catalog(
            mst_item=[{'ItemId': 1, 'ItemName': 'item'}],
            mst_direct_reward=[
                {'DirectRewardGroupId': 5, 'DirectRewardSequenceNo': 2, 'RewardTypeCode': 2, 'RewardTargetId': 1, 'RewardCount': 20},
                {'DirectRewardGroupId': 5, 'DirectRewardSequenceNo': 1, 'RewardTypeCode': 2, 'RewardTargetId': 1, 'RewardCount': 10},
                {'DirectRewardGroupId': 5, 'DirectRewardSequenceNo': 3, 'RewardTypeCode': 2, 'RewardTargetId': 1, 'IsActive': False},
            ],
            mst_present=[{'PresentId': 5, 'PresentSequenceNo': 1, 'RewardTypeCode': 2, 'RewardTargetId': 1, 'RewardCount': 30}],
        ))
        self.assertEqual([10, 20], [r['RewardCount'] for r in resolver.direct(5)])
        self.assertEqual([30], [r['RewardCount'] for r in resolver.present(5)])
        self.assertEqual('mst_present', resolver.present(5)[0]['SourceTable'])
        self.assertEqual([], resolver.direct(0))
        self.assertEqual([], resolver.direct(999))
        self.assertEqual('missing_reward_group', resolver.issues[-1]['Resolution'])

    def test_inactive_target_is_identified_but_not_marked_available(self):
        resolver = RewardResolver(catalog(mst_item=[{'ItemId': 1, 'ItemName': 'historic', 'IsActive': False}]))
        result = resolver.resolve({'RewardTypeCode': 2, 'RewardTargetId': 1})
        self.assertEqual('historic', result['RewardName'])
        self.assertFalse(result['TargetActive'])
        self.assertNotIn('available', result)

    def test_event_export_consumes_honor_and_voice_rewards_and_records_missing_group(self):
        tables = catalog(
            mst_event=[{'EventId': 1, 'EventTitle': 'event', 'EventFormat': 1}],
            mst_event_ranking_reward=[{'EventId': 1, 'PresentId': 8}],
            mst_event_story_section=[{'EventId': 1, 'ReadDirectRewardGroupId': 999}],
            mst_present=[
                {'PresentId': 8, 'PresentSequenceNo': 1, 'RewardTypeCode': 6, 'RewardTargetId': 1, 'RewardCount': 1},
                {'PresentId': 8, 'PresentSequenceNo': 2, 'RewardTypeCode': 10, 'RewardTargetId': 1, 'RewardCount': 1},
            ],
            mst_honor=[{'HonorId': 1, 'HonorName': 'actual honor'}],
            mst_title=[{'TitleId': 1, 'TitleFileName': 'wrong title'}],
            mst_home_voice_product=[{'HomeVoiceProductId': 1, 'DisplayName': 'actual voice'}],
        )
        with tempfile.TemporaryDirectory() as directory:
            context = OutputContext(Path(directory), domain='events')
            with context.activate():
                events.run(session=SimpleNamespace(tables=tables))
            receipt = context.finish()
            self.assertEqual('PASS_WITH_WARNINGS', receipt['status'])
            audit = json.loads((Path(directory) / 'audit_output/event_reward_resolution.json').read_text(encoding='utf-8'))
            self.assertEqual('missing_reward_group', audit['Issues'][0]['Resolution'])
            book = load_workbook(Path(directory) / 'wiki_output/event_archive.xlsx')
            try:
                text = ' '.join(str(c.value) for row in book['RewardSummary'] for c in row)
                self.assertIn('actual honor', text)
                self.assertIn('actual voice', text)
                self.assertNotIn('wrong title', text)
            finally:
                book.close()
