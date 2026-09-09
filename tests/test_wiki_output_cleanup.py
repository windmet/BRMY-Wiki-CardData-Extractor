import json
import os
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from toolkit.domains import audio, events, recipes, snap


@contextmanager
def workspace():
    previous = os.getcwd()
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        (root / 'json_output').mkdir()
        os.chdir(root)
        try:
            yield root
        finally:
            os.chdir(previous)


class WikiOutputCleanupTests(unittest.TestCase):
    def test_event_projection_keeps_editorial_content_and_audits_codes(self):
        event = {
            'EventId': 1, 'Title': 'Event', 'ActivityType': '活动',
            'EventFormat': 3, 'SourceSubtype': 'test', 'SourceSubtypeLabel': 'Branch',
            'StorySections': [{'No': 1, 'Type': 99, 'Title': 'Story'}],
            'Rules': [{'SlideNo': 1, 'Type': 98, 'Description': 'Readable rule'}],
        }
        with workspace() as root:
            (root / 'json_output/Event_Archive.json').write_text(
                json.dumps({'Events': [event]}), encoding='utf-8')
            book = load_workbook(events.export())
            try:
                self.assertNotIn('Mappings', book.sheetnames)
                self.assertEqual(17, book['Overview'].max_column)
                self.assertEqual('Story', book['Story']['E2'].value)
                self.assertEqual('Readable rule', book['Rules']['F2'].value)
                for sheet in book:
                    self.assertFalse(any('类型码' in str(c.value) or '数据来源' in str(c.value)
                                         for c in sheet[1]))
            finally:
                book.close()
            audit = json.loads((root / 'audit_output/event_archive_audit.json').read_text(encoding='utf-8'))
            self.assertEqual(event, audit['Archive']['Events'][0])
            self.assertIn('3', audit['Mappings']['EventFormat'])

    def test_audio_exports_only_text_rows_and_preserves_empty_metadata_in_audit(self):
        with workspace() as root:
            (root / 'voice_403.acb').write_bytes(
                b'binary\x00' + 'title:{ホームボイス①}text:{台詞}'.encode('utf-8')
                + b'\x00title:{empty}text:{}\x00')
            audio.run(str(root))
            for name, first_header in [('card_voice_texts.xlsx', '卡牌ID'),
                                       ('voice_texts.xlsx', '资源文件')]:
                book = load_workbook(root / 'xlsx_output' / name)
                try:
                    sheet = book.active
                    self.assertEqual([first_header, '标题', '日文台词', '中文翻译'],
                                     [c.value for c in sheet[1]])
                    self.assertEqual(2, sheet.max_row)
                    self.assertEqual('台詞', sheet['C2'].value)
                    self.assertIsNone(sheet['D2'].value)
                finally:
                    book.close()
            audit = json.loads((root / 'audit_output/Voice_Text_Index.json').read_text(encoding='utf-8'))
            self.assertEqual(2, len(audit['Entries']))
            self.assertIn('CueIndex', audit['Entries'][0])
            self.assertTrue(any(not item['HasText'] for item in audit['Entries']))

    def test_snap_deduplication_remains_traceable_outside_workbook(self):
        source = {
            'snap_id': 7, 'main_chars': 'Character', 'main_text': 'Text',
            'category': '常驻', 'rarity': 'N', 'scene_raw': 'unknown',
            'comments': [{'char_name': 'Speaker', 'text': 'Comment'}],
        }
        with workspace() as root:
            (root / 'json_output/intermediate_snaps.json').write_text(
                json.dumps([source, {**source, 'snap_id': 9}]), encoding='utf-8'
            )
            snap.export()
            book = load_workbook(root / 'xlsx_output/Snap_Wiki_Data_Clean.xlsx')
            try:
                sheet = book.active
                self.assertEqual(2, sheet.max_row)
                self.assertEqual(14, sheet.max_column)
                self.assertNotIn('折叠重复数', [cell.value for cell in sheet[1]])
                self.assertEqual(7, sheet['A2'].value)
                self.assertEqual('Text', sheet['F2'].value)
                self.assertEqual('Comment', sheet['H2'].value)
            finally:
                book.close()
            audit = json.loads((root / 'audit_output/snap_deduplication.json').read_text(encoding='utf-8'))
            self.assertEqual([{'first_id': 7, 'merged_count': 2, 'source_ids': [7, 9]}], audit)

    def test_multisheet_exports_preserve_locked_files_and_return_real_path(self):
        for module, input_name, output_name, expected_sheets in [
            (events, 'Event_Archive.json', 'event_archive.xlsx', ['Overview', 'Story']),
            (recipes, 'bar_extract.json', 'bar_data_complete.xlsx', ['Ingredients', 'Recipes']),
        ]:
            with self.subTest(domain=module.__name__), workspace() as root:
                (root / 'json_output' / input_name).write_text('{}', encoding='utf-8')
                output = root / 'xlsx_output' / output_name
                output.parent.mkdir()
                output.write_bytes(b'existing editor workbook')
                original_save = Workbook.save

                def save_unless_locked(book, path):
                    if Path(path).resolve() == output:
                        raise PermissionError('simulated Excel lock')
                    return original_save(book, path)

                with patch.object(Workbook, 'save', save_unless_locked):
                    actual = module.export()
                    self.assertEqual(output.with_stem(output.stem + '_new').resolve(), Path(actual).resolve())
                self.assertEqual(b'existing editor workbook', output.read_bytes())
                book = load_workbook(actual)
                try:
                    for sheet_name in expected_sheets:
                        self.assertIn(sheet_name, book.sheetnames)
                        self.assertGreater(book[sheet_name].max_column, 1)
                finally:
                    book.close()


if __name__ == '__main__':
    unittest.main()
