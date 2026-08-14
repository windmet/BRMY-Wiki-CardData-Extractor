import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from toolkit.core.session import MasterDataSession
from toolkit.domains import birthday, items, missions, music, recipes, snap


class NamedDomainTests(unittest.TestCase):
    def test_named_tables_ignore_field_collision_decoys(self):
        names_and_rows = [
            ("decoy_music", [{"MusicId": 999, "DisplayName": "Wrong"}]),
            ("mst_music", [{
                "MusicId": 1,
                "DisplayName": "Right Song",
                "ArtistName": "Right Artist",
                "IsActive": True,
            }]),
            ("decoy_item", [{"ItemId": 999, "ItemTypeCode": 1}]),
            ("mst_item", [{
                "ItemId": 1,
                "ItemName": "Right Item",
                "ItemTypeCode": 1,
                "ItemRarityCode": 1,
                "ItemAttributeCode": 0,
                "ItemDisplayTab": 0,
                "ItemFileName": "item_1",
                "ItemDescription1": "Description",
                "IsActive": True,
            }]),
            ("mst_mission", [{
                "MissionId": 1,
                "Description": "Do # things",
                "MissionType": "sample",
                "IsActive": True,
            }]),
            ("mst_mission_sequence", [{
                "MissionId": 1,
                "MissionSequenceNo": 1,
                "Border": 5,
                "IsHidden": True,
                "IsActive": True,
            }]),
            ("decoy_mission", [{"MissionId": 999, "Description": "Wrong #"}]),
            ("decoy_sequence", [{"MissionId": 999, "IsHidden": True}]),
        ]
        data = [
            {name: [0, 0] for name, _ in names_and_rows},
            *(rows for _, rows in names_and_rows),
        ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            masterdata = root / "master_data.json"
            masterdata.write_text(json.dumps(data), encoding="utf-8")
            session = MasterDataSession.open(masterdata, required_schema={})
            previous = os.getcwd()
            os.chdir(root)
            try:
                music_result = music.extract(session=session)
                items.run(session=session)
                missions.run(session=session)
            finally:
                os.chdir(previous)

            self.assertEqual([1], list(music_result))
            self.assertEqual("Right Song", music_result[1]["DisplayName"])

            item_book = load_workbook(root / "xlsx_output" / "items_catalog.xlsx")
            item_sheet = item_book.active
            self.assertEqual(2, item_sheet.max_row)
            self.assertEqual(1, item_sheet["A2"].value)
            self.assertEqual("Right Item", item_sheet["B2"].value)

            mission_book = load_workbook(root / "xlsx_output" / "hidden_missions.xlsx")
            mission_sheet = mission_book.active
            self.assertEqual(2, mission_sheet.max_row)
            self.assertEqual(1, mission_sheet["A2"].value)
            self.assertEqual("Do 5 things", mission_sheet["F2"].value)

    def test_complex_domains_ignore_field_collision_decoys(self):
        names_and_rows = [
            ("decoy_snap", [{"SnapshotId": 999, "Comment": "Wrong snap"}]),
            ("mst_spin_sticky_note", [{
                "StickyNoteId": 1, "CharacterId": 1, "Comment": "Right note", "IsActive": True,
            }]),
            ("mst_spin_snapshot", [{
                "SnapshotId": 1, "SnapshotCharacterIds": [1], "SpinSetId": 1,
                "Comment": "Right snap", "SnapshotRarityCode": 1,
                "SnapshotSpecialFrameId": 0, "StickyNoteId1": 1, "IsActive": True,
            }]),
            ("mst_spin_set", [{"SpinSetId": 1, "SpinMotionIds": [1], "IsActive": True}]),
            ("mst_spin_motion", [{
                "SpinMotionId": 1, "SpinCharacterMotionIds": [1], "IsActive": True,
            }]),
            ("mst_spin_character_motion", [{
                "SpinCharacterMotionId": 1,
                "SpinCharacterMotionFileName": "scene/right_motion",
                "IsActive": True,
            }]),
            ("decoy_birthday", [{
                "CharacterId": 1, "CharacterBirthdayTextNo": 1,
                "Text": "Wrong birthday", "Year": 2026,
            }]),
            ("mst_character", [{
                "CharacterId": 1, "CharacterNameJpn": "Right Character",
                "CharacterNameEng": "Right Character", "BirthMonth": 6, "BirthDay": 1,
                "IsActive": True,
            }]),
            ("mst_character_collaboration", [{
                "CharacterId": 10001, "CharacterNameJpn": "Right Collaboration",
                "IsActive": True,
            }]),
            ("mst_character_birthday_mini_game_text", [{
                "CharacterId": 1, "CharacterBirthdayTextNo": 1,
                "Text": "Right birthday", "Year": 2026, "IsActive": True,
            }]),
            ("mst_event", [{
                "EventId": 1, "EventTitle": "Right Event", "EventFormat": 1,
                "IsActive": True,
            }]),
            ("mst_event_ingredient", [{
                "IngredientId": 1, "IngredientName": "Right Ingredient",
                "IngredientDescription": "Right ingredient description",
                "IngredientFileName": "right_ingredient", "IsActive": True,
            }]),
            ("mst_event_menu", [{
                "ShiftId": 1, "MenuSequenceNo": 1, "RecipeId": 1, "IsActive": True,
            }]),
            ("mst_event_recipe", [{
                "RecipeId": 1, "RecipeName": "Right Recipe", "IngredientIds": [1],
                "EventId": 1, "RecommendCharacterId": 1, "IsActive": True,
            }]),
            ("decoy_recipe", [{
                "RecipeId": 999, "RecipeName": "Wrong Recipe", "IngredientIds": [1],
            }]),
        ]
        data = [
            {name: [0, 0] for name, _ in names_and_rows},
            *(rows for _, rows in names_and_rows),
        ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            masterdata = root / "master_data.json"
            masterdata.write_text(json.dumps(data), encoding="utf-8")
            session = MasterDataSession.open(masterdata, required_schema={})
            previous = os.getcwd()
            os.chdir(root)
            try:
                snaps = snap.extract(session=session)
                characters, birthday_texts = birthday.extract(session=session)
                _, recipe_characters, ingredients, recipe_rows, shift_menus = recipes.extract(session=session)
            finally:
                os.chdir(previous)

        self.assertEqual([1], [row["snap_id"] for row in snaps])
        self.assertEqual("Right snap", snaps[0]["main_text"])
        self.assertEqual("right_motion", snaps[0]["scene_raw"])
        self.assertEqual("Right Character", characters[1]["Name"])
        self.assertEqual("Right birthday", birthday_texts[1][1])
        self.assertEqual("Right Collaboration", recipe_characters[10001])
        self.assertEqual([1], list(ingredients))
        self.assertEqual([1], list(recipe_rows))
        self.assertEqual([1], list(shift_menus))


if __name__ == "__main__":
    unittest.main()
