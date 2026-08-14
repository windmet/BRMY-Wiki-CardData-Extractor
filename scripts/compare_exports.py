#!/usr/bin/env python3
"""Compare two toolkit export directories by data content, not file bytes."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def compare_json(expected, actual):
    left = json.loads(expected.read_text(encoding="utf-8"))
    right = json.loads(actual.read_text(encoding="utf-8"))
    return None if left == right else "JSON objects differ"


def compare_xlsx(expected, actual):
    left = load_workbook(expected, data_only=False, read_only=False)
    right = load_workbook(actual, data_only=False, read_only=False)
    if left.sheetnames != right.sheetnames:
        return f"sheet names differ: {left.sheetnames!r} != {right.sheetnames!r}"
    for title in left.sheetnames:
        left_sheet = left[title]
        right_sheet = right[title]
        left_size = (left_sheet.max_row, left_sheet.max_column)
        right_size = (right_sheet.max_row, right_sheet.max_column)
        if left_size != right_size:
            return f"{title}: dimensions differ: {left_size} != {right_size}"
        left_merged = sorted(map(str, left_sheet.merged_cells.ranges))
        right_merged = sorted(map(str, right_sheet.merged_cells.ranges))
        if left_merged != right_merged:
            return f"{title}: merged ranges differ"
        for row in left_sheet.iter_rows():
            for left_cell in row:
                right_value = right_sheet[left_cell.coordinate].value
                if left_cell.value != right_value:
                    return (
                        f"{title}!{left_cell.coordinate}: "
                        f"{left_cell.value!r} != {right_value!r}"
                    )
                right_cell = right_sheet[left_cell.coordinate]
                if left_cell.data_type != right_cell.data_type:
                    return (
                        f"{title}!{left_cell.coordinate}: data types differ: "
                        f"{left_cell.data_type!r} != {right_cell.data_type!r}"
                    )
                if left_cell.number_format != right_cell.number_format:
                    return (
                        f"{title}!{left_cell.coordinate}: number formats differ: "
                        f"{left_cell.number_format!r} != {right_cell.number_format!r}"
                    )
    return None


def files_in(root, directory, suffix):
    path = root / directory
    return {item.name: item for item in path.glob(f"*{suffix}")} if path.is_dir() else {}


def compare_exports(expected_root, actual_root):
    groups = [
        ("json_output", ".json", compare_json),
        ("xlsx_output", ".xlsx", compare_xlsx),
    ]
    failures = []
    for directory, suffix, comparator in groups:
        expected = files_in(expected_root, directory, suffix)
        actual = files_in(actual_root, directory, suffix)
        if expected.keys() != actual.keys():
            missing = sorted(expected.keys() - actual.keys())
            extra = sorted(actual.keys() - expected.keys())
            failures.append(f"{directory}: missing={missing}, extra={extra}")
            continue
        for name in sorted(expected):
            difference = comparator(expected[name], actual[name])
            status = "PASS" if difference is None else "FAIL"
            print(f"[{status}] {directory}/{name}")
            if difference:
                failures.append(f"{directory}/{name}: {difference}")
    return failures


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("expected", type=Path, help="previous accepted export root")
    parser.add_argument("actual", type=Path, help="new export root from the same inputs")
    args = parser.parse_args()

    failures = compare_exports(args.expected.resolve(), args.actual.resolve())
    if failures:
        print("\nExport regression failed:")
        for failure in failures:
            print(f"- {failure}")
        raise SystemExit(1)
    print(
        "\nExport regression passed: all JSON objects and XLSX cell values, "
        "types, and number formats match."
    )


if __name__ == "__main__":
    main()
