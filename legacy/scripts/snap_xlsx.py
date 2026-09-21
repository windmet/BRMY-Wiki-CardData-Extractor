import openpyxl
from collections import defaultdict
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

INPUT_XLSX = 'Snap_Wiki_Data_Clean.xlsx'
OUTPUT_XLSX = 'Snap_Wiki_For_Online_Doc.xlsx'

# 全角色去姓留名对照表
NAME_MAP = {
    "皇坂逢": "逢", "城瀬由鶴": "由鶴", "須王芦佳": "芦佳", "綾戸恋": "恋", "宇京真央": "真央",
    "樋宮明星": "明星", "環野揺": "揺", "槻本大河": "大河", "壱川春日": "春日", "隠岐谷誓": "誓",
    "節見静": "静", "御門尊": "尊", "新開戦": "戦", "相沢篠信": "篠信", "在間樹帆": "樹帆",
    "祠堂恭耶": "恭耶", "立科吏来": "吏来", "恩田灯世": "灯世", "新名有": "有", "神家": "神家",
    "麻波麗": "麗"
}

def split_surname_name(full_name):
    """拆分姓和名"""
    if full_name == "神家": return "", "神家"
    first_name = NAME_MAP.get(full_name, full_name)
    surname = full_name.replace(first_name, "")
    return surname, first_name

def format_cell_text(row_dict):
    """整理主文案和评论为换行文本 (Excel换行效果)"""
    text = str(row_dict.get('主文案', '')).strip()
    text = text.replace('<br>', '\n').replace('<br/>', '\n')
    
    comments = []
    for i in range(1, 5):
        c_name = row_dict.get(f'评论{i}_角色')
        c_text = row_dict.get(f'评论{i}_文案')
        if c_name and str(c_name).strip():
            short_name = NAME_MAP.get(c_name, c_name)
            comments.append(f"{short_name}：{c_text}")
            
    if comments:
        if text:
            text += "\n" + "\n".join(comments)
        else:
            text = "\n".join(comments)
            
    return text

def style_ws(ws):
    """应用美观的排版样式"""
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    header_font = Font(name='微软雅黑', size=11, bold=True, color='333333')
    header_fill = PatternFill(start_color='F2F4F8', end_color='F2F4F8', fill_type='solid') # 淡蓝灰色
    data_font = Font(name='微软雅黑', size=10, color='333333')
    
    align_center = Alignment(vertical='center', horizontal='center', wrap_text=True)
    align_left_top = Alignment(vertical='top', horizontal='left', wrap_text=True)

    # 样式化所有单元格
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
            if cell.row in [1, 2]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center
            else:
                cell.font = data_font
                cell.alignment = align_left_top

    # 自动设定行高
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20
    for r in range(3, ws.max_row + 1):
        ws.row_dimensions[r].height = 55 # 预设舒适的多行高度

    # 自动调整列宽
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col[0].column in [1, 2] and ws.title in ["单人通常", "生日限定"]:
            ws.column_dimensions[col_letter].width = 12
        elif col[0].column == 1 and ws.title in ["双人通常", "多人通常"]:
            ws.column_dimensions[col_letter].width = 25
        else:
            ws.column_dimensions[col_letter].width = 38 # 给Spot留足宽度防止挤压

def write_sheet_data(ws, data_dict, is_single=False, spot_prefix="通常Spot", specified_scenes=None):
    """矩阵拼装与双层表头写入"""
    if not data_dict:
        return

    # 1. 确定所有的场景列（普通场景按数量，生日场景按特定名称对齐）
    if specified_scenes:
        scenes_list = sorted(list(specified_scenes))
        max_scenes = len(scenes_list)
    else:
        max_scenes = max([len(scenes) for scenes in data_dict.values()])
        scenes_list = [f"{spot_prefix} {i}" for i in range(1, max_scenes + 1)]

    # 2. 写入双层表头
    start_col = 3 if is_single else 2
    
    # 第一行表头
    if is_single:
        ws.cell(row=1, column=1, value="姓")
        ws.cell(row=1, column=2, value="名")
        ws.merge_cells('A1:A2')
        ws.merge_cells('B1:B2')
    else:
        ws.cell(row=1, column=1, value="相片主角")
        ws.merge_cells('A1:A2')

    for i, scene_title in enumerate(scenes_list):
        col_idx = start_col + i * 3
        ws.cell(row=1, column=col_idx, value=scene_title)
        
        # 合并三列 (N, R, SR)
        start_letter = get_column_letter(col_idx)
        end_letter = get_column_letter(col_idx + 2)
        ws.merge_cells(f"{start_letter}1:{end_letter}1")

    # 第二行表头
    for i in range(max_scenes):
        col_idx = start_col + i * 3
        ws.cell(row=2, column=col_idx, value="N")
        ws.cell(row=2, column=col_idx + 1, value="R")
        ws.cell(row=2, column=col_idx + 2, value="SR")

    # 3. 填充数据
    for protagonist in sorted(data_dict.keys()):
        scenes = data_dict[protagonist]
        row_data = []
        if is_single:
            surname, first_name = split_surname_name(protagonist)
            row_data.extend([surname, first_name])
        else:
            row_data.append(protagonist)

        # 映射填充
        if specified_scenes:
            # 针对生日：按特定场景名（如二轮、三轮）放入对应列
            for scene_title in scenes_list:
                rarities = scenes.get(scene_title, {})
                row_data.extend([rarities.get('N', ''), rarities.get('R', ''), rarities.get('SR', '')])
        else:
            # 针对常驻：按顺序依次塞入 Spot 1, Spot 2...
            sorted_scenes = sorted(scenes.keys())
            for idx in range(max_scenes):
                if idx < len(sorted_scenes):
                    scene_name = sorted_scenes[idx]
                    rarities = scenes[scene_name]
                    row_data.extend([rarities.get('N', ''), rarities.get('R', ''), rarities.get('SR', '')])
                else:
                    row_data.extend(['', '', ''])

        ws.append(row_data)

    # 4. 样式美化
    style_ws(ws)

def main():
    print(f"[*] 正在读取已清洗的数据表: {INPUT_XLSX} ...")
    try:
        wb_in = openpyxl.load_workbook(INPUT_XLSX)
        ws_in = wb_in.active
    except Exception as e:
        print(f"[!] 读取失败，请检查文件: {e}")
        return

    headers = [cell.value for cell in ws_in[1]]
    
    # 初始化四种数据池
    normal_data = {
        'single': defaultdict(lambda: defaultdict(dict)),
        'pair': defaultdict(lambda: defaultdict(dict)),
        'group': defaultdict(lambda: defaultdict(dict))
    }
    birthday_data = defaultdict(lambda: defaultdict(dict))
    all_birthday_scenes = set()

    print("[*] 正在解析数据并执行角色合并...")
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        
        protagonist = str(row_dict.get('相片主角', '')).strip()
        category = row_dict.get('所属分类', '')
        rarity = row_dict.get('稀有度', '')
        scene = row_dict.get('互动场景', '')
        
        cell_text = format_cell_text(row_dict)

        if "生日" in category or "BirthDay" in cell_text or "Birthday" in cell_text:
            birthday_data[protagonist][scene][rarity] = cell_text
            all_birthday_scenes.add(scene)
        else:
            ampersand_count = protagonist.count('&')
            if ampersand_count == 0:
                normal_data['single'][protagonist][scene][rarity] = cell_text
            elif ampersand_count == 1:
                normal_data['pair'][protagonist][scene][rarity] = cell_text
            else:
                normal_data['group'][protagonist][scene][rarity] = cell_text

    # 创建新的 Wiki 专用 Excel
    wb_out = openpyxl.Workbook()
    
    # 1. 单人通常
    ws1 = wb_out.active
    ws1.title = "单人通常"
    write_sheet_data(ws1, normal_data['single'], is_single=True, spot_prefix="通常Spot")

    # 2. 双人通常
    ws2 = wb_out.create_sheet(title="双人通常")
    write_sheet_data(ws2, normal_data['pair'], is_single=False, spot_prefix="双人Spot")

    # 3. 多人通常
    ws3 = wb_out.create_sheet(title="多人通常")
    write_sheet_data(ws3, normal_data['group'], is_single=False, spot_prefix="组/部门Spot")

    # 4. 生日限定
    ws4 = wb_out.create_sheet(title="生日限定")
    write_sheet_data(ws4, birthday_data, is_single=True, specified_scenes=all_birthday_scenes)

    # 保存
    wb_out.save(OUTPUT_XLSX)
    print(f"[+] Excel 生成成功！已保存至: {OUTPUT_XLSX}")
    print("[+] 你现在可以直接用 Excel/WPS 打开它，完美的排版已经呈现！")

if __name__ == '__main__':
    main()