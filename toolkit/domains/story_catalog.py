"""Story metadata index; never guesses script filenames or parses story bodies."""
from collections import defaultdict
import json

from openpyxl import Workbook

from ..core.availability import assessment_time, date_window
from ..core.rewards import RewardResolver
from ..core.scanner import save_json
from ..core.exporter import audit_path, xlsx_path, save_workbook_safely
from ..core.output import record_warning


SPECS = [
    ('main', 'mst_main_story_section', ('MainStoryThreadNo', 'MainStoryChapterNo', 'MainStorySectionNo'), 'mst_main_story_chapter', ('MainStoryThreadNo', 'MainStoryChapterNo')),
    ('character', 'mst_character_story_section', ('CharacterId', 'CharacterStoryChapterNo', 'CharacterStorySectionNo'), 'mst_character_story', ('CharacterId', 'CharacterStoryChapterNo')),
    ('card', 'mst_character_card_story_section', ('CharacterCardId', 'CharacterCardStorySectionNo'), 'mst_character_card_story', ('CharacterCardId',)),
    ('event', 'mst_event_story_section', ('EventId', 'EventStorySectionNo'), 'mst_event_story', ('EventId',)),
    ('login', 'mst_login_story_section', ('LoginStoryId', 'LoginStorySectionNo'), None, ()),
    ('puzzle', 'mst_puzzle_story', ('StoryTypeCode', 'StoryTargetBaseId', 'StoryTargetChapterId', 'StoryTargetSectionNo', 'PuzzleStoryNo'), 'mst_puzzle_story_stage', ('StoryTypeCode', 'StoryTargetBaseId', 'StoryTargetChapterId', 'StoryTargetSectionNo')),
    ('main', 'mst_main_story_from_other_story_type_section', ('MainStoryThreadNo', 'MainStoryFromOtherStoryTypeChapterNo', 'MainStoryFromOtherStoryTypeSectionNo'), 'mst_main_story_from_other_story_type_chapter', ('MainStoryThreadNo', 'MainStoryFromOtherStoryTypeChapterNo')),
]


def extract(session, *, as_of=None):
    tables = session.tables; now = assessment_time(as_of); rewards = RewardResolver(tables)
    sections, issues = [], []
    for kind, table, keys, parent_table, parent_keys in SPECS:
        parents = defaultdict(list)
        if parent_table:
            for row in tables.rows(parent_table, active_only=True):
                parents[tuple(row.get(k) for k in parent_keys)].append(row)
        grouped = defaultdict(list)
        for row in tables.rows(table, active_only=True):
            grouped[tuple(row.get(k) for k in keys)].append(row)
        for key, rows in sorted(grouped.items()):
            if len(rows) != 1 or None in key:
                issues.append({'Status': 'ambiguous_or_missing_key', 'Table': table, 'Key': key, 'Rows': rows})
                continue
            row = rows[0]; parent = parents.get(tuple(row.get(k) for k in parent_keys), []) if parent_table else []
            if parent_table and len(parent) != 1:
                issues.append({'Status': 'missing_or_ambiguous_parent', 'Table': table, 'Key': key, 'ParentTable': parent_table})
            release = row.get('ReleaseDateTime') or row.get('StartTime')
            if not release and len(parent) == 1:
                release = parent[0].get('ReleaseDateTime')
            # A release date has no stated expiry; login has an explicit window.
            end = row.get('EndTime') if kind == 'login' else '9999-01-01T00:00:00+00:00'
            scripts = {k: v for source in [row, *parent] for k, v in source.items()
                       if k in ('StoryFileName', 'PuzzleClearStoryFileName', 'RereadStoryFileName') and v}
            sections.append({'Kind': kind, 'SourceTable': table, 'Key': list(key),
                'Title': row.get('TitleName', ''), 'ParentTable': parent_table, 'Parents': parent,
                'Raw': row, 'ReadRewards': rewards.direct(row.get('ReadDirectRewardGroupId')),
                'Availability': date_window(release, end, as_of=now),
                'ExplicitScripts': scripts, 'ScriptResolution': 'explicit' if scripts else 'not_declared'})
    return {'Sections': sections, 'Issues': issues, 'RewardIssues': rewards.issues,
            'AsOf': now.isoformat(), 'Timezone': 'UTC',
            'RawTables': {n: tables.rows(n) for n in tables.names if 'story' in n}}


def export(data):
    wb = Workbook(); ws = wb.active; ws.title = 'StoryCatalog'
    ws.append(['类别', '复合编号', '章节标题', '上级标题', '有语音标志', '发布日期/开始', '结束',
               '日期状态', '阅读奖励', '明确脚本名', '脚本关联状态'])
    labels = {'main': '主线', 'character': '角色', 'card': '卡牌', 'event': '活动', 'login': '登录', 'puzzle': '剧情内Puzzle'}
    for entry in data['Sections']:
        raw = entry['Raw']
        ws.append([labels[entry['Kind']], entry['SourceTable'] + ':' + '/'.join(map(str, entry['Key'])),
                   entry['Title'], ' / '.join(p.get('TitleName', '') for p in entry['Parents']),
                   raw.get('HasVoiceFile'), entry['Availability']['RawStart'], raw.get('EndTime'),
                   entry['Availability']['Status'], ' / '.join(f"{r['RewardName']} x{r['RewardCount']}" for r in entry['ReadRewards']),
                   ' / '.join(entry['ExplicitScripts'].values()), entry['ScriptResolution']])
    conditions = wb.create_sheet('Requirements'); conditions.append(['类别', '复合编号', '门槛字段', '原值'])
    for entry in data['Sections']:
        for key, value in entry['Raw'].items():
            if key.startswith('Key'):
                conditions.append([labels[entry['Kind']], '/'.join(map(str, entry['Key'])), key, json.dumps(value, ensure_ascii=False)])
    notes = wb.create_sheet('Notes'); notes.append(['核对时刻 UTC', data['AsOf']])
    notes.append(['用途', '剧情元数据索引；不等于正文已解析或账号已解锁'])
    notes.append(['文件名', '仅列原表明确声明的脚本，不按编号拼造'])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or '')) for c in column[:80]) + 2))
    return save_workbook_safely(wb, xlsx_path('story_catalog.xlsx'))


def run(session=None, *, as_of=None):
    data = extract(session, as_of=as_of)
    save_json(data, audit_path('story_catalog.json'))
    if data['Issues'] or data['RewardIssues']:
        record_warning(f"剧情索引有 {len(data['Issues'])} 条关系异常、{len(data['RewardIssues'])} 条奖励缺项，详见 story_catalog.json")
    return export(data)
