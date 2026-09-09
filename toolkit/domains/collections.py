"""Collection identities, card relations and typed reward references."""
from openpyxl import Workbook

from ..core.rewards import REWARD_TARGETS, reward_reference_index
from ..core.availability import assessment_time, date_window
from ..core.exporter import audit_path, xlsx_path, save_workbook_safely
from ..core.scanner import save_json
from ..core.output import record_warning


COLLECTION_TYPES = (4, 5, 6, 7, 3, 2)


def extract(session, *, as_of=None):
    tables = session.tables; now = assessment_time(as_of)
    refs = reward_reference_index(tables)
    cards = tables.group_by('mst_character_card', 'CharacterCardId', active_only=False, required=False)
    records, issues = [], []
    for code in COLLECTION_TYPES:
        kind, label, table, key, name_key = REWARD_TARGETS[code]
        for raw in tables.rows(table, active_only=True):
            identifier = raw.get(key); name = raw.get(name_key)
            card_id = raw.get('CharacterCardId'); card_matches = cards.get(card_id, []) if card_id else []
            if card_id and len(card_matches) != 1:
                issues.append({'Kind': kind, 'Id': identifier, 'Status': 'missing_or_ambiguous_card', 'CardId': card_id})
            records.append({'Kind': kind, 'Label': label, 'RewardTypeCode': code, 'Id': identifier,
                'Name': name or f'{label}({identifier})', 'NameStatus': 'named' if name else 'name_unavailable',
                'Raw': raw, 'Cards': card_matches, 'RewardReferences': refs.get((code, identifier), []),
                'Availability': date_window(raw.get('ReleaseDateTime') or raw.get('StartTime'),
                    raw.get('EndTime') or ('9999-01-01T00:00:00Z' if raw.get('ReleaseDateTime') else None), as_of=now),
                'AcquisitionEntryStatus': 'not_resolved'})
    return {'Collections': records, 'Issues': issues, 'AsOf': now.isoformat(), 'Timezone': 'UTC'}


def export(data):
    wb = Workbook(); wb.remove(wb.active)
    for code in COLLECTION_TYPES:
        kind, label, *_ = REWARD_TARGETS[code]
        sheet = wb.create_sheet(kind)
        sheet.append(['编号', '名称', '名称状态', '角色ID', '关联卡牌', '奖励引用数', '日期状态'])
        for entry in data['Collections']:
            if entry['Kind'] != kind:
                continue
            sheet.append([entry['Id'], entry['Name'], entry['NameStatus'], entry['Raw'].get('CharacterId'),
                          ' / '.join(f"{c['CharacterCardId']}: {c.get('CharacterCardName', '')}" for c in entry['Cards']),
                          len(entry['RewardReferences']), entry['Availability']['Status']])
    notes = wb.create_sheet('Notes'); notes.append(['核对时刻 UTC', data['AsOf']])
    notes.append(['获取关系', '卡牌关联、奖励组引用与最终获取入口分开；本批最终入口尚未解析'])
    notes.append(['名称', '空称号/徽章名保留编号，不以图标文件名代替'])
    notes.append(['日期', '日期区间不代表账号已拥有或解锁；无日期时保留未知'])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or '')) for c in column[:80]) + 2))
    return save_workbook_safely(wb, xlsx_path('collections.xlsx'))


def run(session=None, *, as_of=None):
    data = extract(session, as_of=as_of)
    save_json(data, audit_path('collections.json'))
    if data['Issues']:
        record_warning(f"收藏档案有 {len(data['Issues'])} 条关系缺项，详见 collections.json")
    return export(data)
