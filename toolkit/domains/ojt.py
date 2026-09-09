"""OJT masterdata archive; chart coordinate files remain a separate input."""
from openpyxl import Workbook

from ..core.availability import assessment_time, date_window, STATUS_LABELS
from ..core.exporter import audit_path, xlsx_path, save_workbook_safely
from ..core.scanner import save_json
from ..core.rewards import RewardResolver
from ..core.output import record_warning


def extract(session, *, as_of=None):
    tables = session.tables
    now = assessment_time(as_of)
    rewards = RewardResolver(tables)
    issues = []

    def grouped(name, key):
        return tables.group_by(name, key, required=False)

    events = grouped('mst_event', 'EventId')
    charts = grouped('mst_event_ojt_chart', 'EventId')
    shifts = grouped('mst_event_ojt_shift', 'EventId')
    boxes = grouped('mst_event_ojt_prize_box', 'OjtShiftId')
    stages = grouped('mst_event_ojt_puzzle_stage', 'OjtShiftId')
    training = grouped('mst_event_ojt_training_reward', 'OjtShiftId')
    texts = grouped('mst_event_ojt_terminal_character_text', 'OjtShiftId')
    fixed = grouped('mst_event_ojt_prize_box_reward', 'PrizeBoxRewardId')
    random = grouped('mst_event_ojt_prize_box_reward_random', 'PrizeBoxRewardRandomId')

    def expand(groups, key, table, location):
        if key in (None, 0):
            return []
        rows = groups.get(key, [])
        if not rows:
            issues.append(dict(location, Status='missing_reward_group', Table=table, Id=key))
        return [dict(rewards.resolve(row), SourceTable=table) for row in rows]

    archive = []
    for raw in tables.require('mst_event_c', active_only=True):
        eid = raw['EventId']
        matches = events.get(eid, [])
        if len(matches) != 1 or matches[0].get('EventFormat') != 5:
            issues.append({'EventId': eid, 'Status': 'invalid_event_owner', 'Raw': raw})
            continue
        event = matches[0]
        entry = {'EventId': eid, 'Title': event.get('EventTitle', ''), 'RawEvent': event,
                 'RawOjt': raw, 'Charts': charts.get(eid, []), 'Shifts': []}
        if len(entry['Charts']) != 1:
            issues.append({'EventId': eid, 'Status': 'missing_or_ambiguous_chart'})
        for shift in sorted(shifts.get(eid, []), key=lambda r: r['OjtShiftId']):
            sid = shift['OjtShiftId']
            item = {'RawShift': shift, 'Stages': stages.get(sid, []),
                    'TerminalTexts': texts.get(sid, []), 'TrainingRewards': [], 'PrizeBoxes': [],
                    'Availability': date_window(shift.get('StartTime'), shift.get('EndTime'), as_of=now)}
            for row in training.get(sid, []):
                item['TrainingRewards'].append({'Raw': row, 'Rewards': rewards.direct(row.get('DirectRewardGroupId'))})
            for box in sorted(boxes.get(sid, []), key=lambda r: r['PrizeBoxNo']):
                loc = {'EventId': eid, 'OjtShiftId': sid, 'PrizeBoxNo': box['PrizeBoxNo']}
                item['PrizeBoxes'].append({'Raw': box,
                    'FixedRewards': expand(fixed, box.get('PrizeBoxRewardId'), 'mst_event_ojt_prize_box_reward', loc),
                    'RandomRewards': expand(random, box.get('PrizeBoxRewardRandomId'), 'mst_event_ojt_prize_box_reward_random', loc)})
            entry['Shifts'].append(item)
        archive.append(entry)
    used_shifts = {s['RawShift']['OjtShiftId'] for e in archive for s in e['Shifts']}
    for name, groups in [('PrizeBoxes', boxes), ('Stages', stages), ('TrainingRewards', training), ('TerminalTexts', texts)]:
        for sid, rows in groups.items():
            if sid not in used_shifts:
                issues.append({'Status': 'unlinked_shift', 'Section': name, 'OjtShiftId': sid, 'Rows': rows})
    return {'Events': archive, 'AsOf': now.isoformat(), 'Timezone': 'UTC',
            'ChartCoordinates': 'not_loaded', 'Issues': issues, 'RewardIssues': rewards.issues,
            'RawTables': {name: tables.rows(name) for name in tables.names if name.startswith('mst_event_ojt_')}}


def export(data):
    wb = Workbook()
    wb.remove(wb.active)
    headers = {
        'Charts': ['活动ID', '活动名', '题面', '左', '右', '下', '上', '坐标资源名'],
        'Shifts': ['活动ID', '轮次ID', '角色ID', '开始', '结束', '日期状态', '训练开始脚本', '训练结束脚本'],
        'Stages': ['活动ID', '轮次ID', 'Phase原值', '关卡序号', '地图ID', 'BREAK最低分', 'BREAK最高分', '租借卡牌ID'],
        'TerminalTexts': ['活动ID', '轮次ID', '文本类型原值', '文本', '播放时长原值'],
        'TrainingRewards': ['活动ID', '轮次ID', 'Phase原值', '奖励'],
        'PrizeBoxes': ['活动ID', '轮次ID', '箱号', '消耗奖牌', '固定池权重原值', '随机池权重原值'],
        'BoxRewards': ['活动ID', '轮次ID', '箱号', '奖励池', '奖励', '数量', '库存', '重点奖励', '池内权重原值'],
    }
    sheets = {name: wb.create_sheet(name) for name in headers}
    for name, cols in headers.items():
        sheets[name].append(cols)
    for event in data['Events']:
        eid = event['EventId']
        for c in event['Charts']:
            sheets['Charts'].append([eid, event['Title'], *[c.get(k, '') for k in ('ChartTitle', 'LeftString', 'RightString', 'BottomString', 'TopString', 'ChartFileName')]])
        for s in event['Shifts']:
            raw = s['RawShift']; sid = raw['OjtShiftId']; prefix = [eid, sid]
            sheets['Shifts'].append(prefix + [raw.get('CharacterId'), raw.get('StartTime'), raw.get('EndTime'),
                STATUS_LABELS[s['Availability']['Status']], raw.get('TrainingPuzzleStartStoryFileName'), raw.get('TrainingPuzzleClearStoryFileName')])
            for row in s['Stages']:
                sheets['Stages'].append(prefix + [row.get(k) for k in ('OjtPhase', 'PuzzleStageNo', 'PuzzleMapId', 'BreakScoreMin', 'BreakScoreMax')] + [', '.join(map(str, row.get('RentalCharacterCardIds', [])))])
            for row in s['TerminalTexts']:
                sheets['TerminalTexts'].append(prefix + [row.get(k) for k in ('OjtTerminalCharacterTextType', 'Text', 'TextPlayTime')])
            for row in s['TrainingRewards']:
                sheets['TrainingRewards'].append(prefix + [row['Raw'].get('OjtPhase'), ' / '.join(f"{r['RewardName']} x{r['RewardCount']}" for r in row['Rewards'])])
            for box in s['PrizeBoxes']:
                raw = box['Raw']; bp = prefix + [raw['PrizeBoxNo']]
                sheets['PrizeBoxes'].append(bp + [raw.get(k) for k in ('CostPrizeMedalCount', 'RewardLotteryRate', 'RewardRandomLotteryRate')])
                for key, label in [('FixedRewards', '固定池'), ('RandomRewards', '随机池')]:
                    for reward in box[key]:
                        rr = reward['RawReward']
                        sheets['BoxRewards'].append(bp + [label, reward['RewardName'], reward['RewardCount'], rr.get('PrizeStock'), rr.get('IsPickUp'), rr.get('LotteryRate')])
    notes = wb.create_sheet('Notes')
    notes.append(['核对时刻 UTC', data['AsOf']])
    notes.append(['日期状态', '日期区间内不代表账号已解锁'])
    notes.append(['坐标', '未加载 .s2bchart；当前仅题面与四轴文案'])
    notes.append(['权重', '保留原值，未推算最终概率'])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or '')) for c in column[:80]) + 2))
    return save_workbook_safely(wb, xlsx_path('ojt_archive.xlsx'))


def run(session=None, *, as_of=None):
    data = extract(session, as_of=as_of)
    save_json(data, audit_path('ojt_archive.json'))
    if data['Issues'] or data['RewardIssues']:
        record_warning(f"OJT 存在 {len(data['Issues'])} 条关系异常、{len(data['RewardIssues'])} 条奖励缺项，详见 ojt_archive.json")
    return export(data)
