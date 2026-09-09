"""Event archive extraction and export."""
from ..core.scanner import load_json, save_json
from ..core.exporter import json_path, xlsx_path, save_workbook_safely, audit_path
from ..core.data import clean_text
from ..core.tables import TableCatalog
from ..core.rewards import RewardResolver, REWARD_TYPE_MAP
from ..core.output import record_warning
from .event_classification import classify_event
from .event_missions import extract_event_missions
from ..core.availability import assessment_time, date_window, STATUS_LABELS

INPUT_JSON = 'master_data.json'

EVENT_FORMAT_MAP = {
    1: {"Key": "standard_exchange", "Label": "普通点数/交换所活动"},
    2: {"Key": "story_campaign_a", "Label": "活动A类/剧情活动"},
    3: {"Key": "bar_sales", "Label": "酒保营业活动"},
    4: {"Key": "special_bar_sales", "Label": "特殊营业活动"},
    5: {"Key": "ojt", "Label": "OJT活动"},
    6: {"Key": "accumulate_item", "Label": "累积道具活动"},
}

SOURCE_SUBTYPE_MAP = {
    "event_a": "mst_event_a：普通活动扩展表",
    "event_b": "mst_event_b：营业/酒保活动扩展表",
    "event_c": "mst_event_c：OJT/特殊活动扩展表",
    "accumulate_item": "mst_event_accumulate_item：累积道具活动扩展表",
    "ojt": "mst_event_ojt_shift：OJT轮次表",
    "unknown": "未匹配到活动扩展表",
}

RARITY_MAP = {1: "R", 2: "SR", 3: "SSR", 101: "XR", 102: "CR"}
ATTRIBUTE_MAP = {1: "Sun", 2: "Moon", 3: "Star"}



def _names(ids, mapping):
    return [mapping.get(i, f"Unknown({i})") for i in ids or []]


def _card_summary(card, character_map):
    if not card:
        return ""
    char_name = character_map.get(card.get("CharacterId"), f"Char_{card.get('CharacterId')}")
    rarity = RARITY_MAP.get(card.get("CardRarityCode"), card.get("CardRarityCode"))
    attr = ATTRIBUTE_MAP.get(card.get("CardAttributeCode"), card.get("CardAttributeCode"))
    return f"{card.get('CharacterCardName', '')} / {char_name} / {rarity} / {attr}"


def extract(session=None, *, as_of=None):
    as_of = assessment_time(as_of)
    tables = session.tables if session else TableCatalog(load_json(INPUT_JSON))

    events = tables.require("mst_event", active_only=True)

    def optional_by_id(name, key):
        return tables.by_id(name, key, required=False)

    def optional_group(name, key):
        return tables.group_by(name, key, required=False)

    event_a = optional_by_id("mst_event_a", "EventId")
    event_b = optional_by_id("mst_event_b", "EventId")
    event_c = optional_by_id("mst_event_c", "EventId")
    event_acc = optional_by_id("mst_event_accumulate_item", "EventId")
    event_story = optional_by_id("mst_event_story", "EventId")
    stories = optional_group("mst_event_story_section", "EventId")
    rules = optional_group("mst_event_rule_window", "EventId")
    shifts = optional_group("mst_event_shift", "EventId")
    ojt_shifts = optional_group("mst_event_ojt_shift", "EventId")
    special_by_shift = optional_group("mst_event_special_time", "ShiftId")
    recipes = optional_group("mst_event_recipe", "EventId")
    accumulate_rewards = optional_group("mst_event_accumulate_item_reward", "EventId")
    sales_rewards = optional_group("mst_event_sales_reward", "EventId")
    ranking_rewards = optional_group("mst_event_ranking_reward", "EventId")
    ingredient_stages = optional_group("mst_event_puzzle_stage_ingredient", "EventId")

    characters = optional_by_id("mst_character", "CharacterId")
    cards = optional_by_id("mst_character_card", "CharacterCardId")
    items = optional_by_id("mst_item", "ItemId")
    ingredients = optional_by_id("mst_event_ingredient", "IngredientId")
    character_map = {k: v.get("CharacterNameJpn", "") for k, v in characters.items()}
    card_map = {k: _card_summary(v, character_map) for k, v in cards.items()}
    item_map = {k: v.get("ItemName", "") for k, v in items.items()}
    ingredient_map = {k: v.get("IngredientName", "") for k, v in ingredients.items()}
    rewards = RewardResolver(tables, card_names=card_map)
    event_missions, mission_audit = extract_event_missions(tables)
    for missions in event_missions.values():
        for mission in missions:
            raw = mission['RawMission']
            mission['Availability'] = date_window(raw.get('StartTime'), raw.get('EndTime'),
                                                  as_of=as_of, active=raw.get('IsActive', True))

    archive = []
    for event in sorted(events, key=lambda e: e.get("EventId", 0)):
        event_id = event.get("EventId")
        subtype_rows = [r for r in (event_a.get(event_id), event_b.get(event_id), event_c.get(event_id), event_acc.get(event_id)) if r]
        subtype = (
            "event_a" if event_id in event_a else
            "event_b" if event_id in event_b else
            "event_c" if event_id in event_c else
            "accumulate_item" if event_id in event_acc else
            "ojt" if event_id in ojt_shifts else
            "unknown"
        )
        main_sub = subtype_rows[0] if subtype_rows else {}
        character_ids = main_sub.get("CharacterIds", [])
        card_ids = main_sub.get("CharacterCardIds", []) or main_sub.get("PickUpCharacterCardIds", [])
        pickup_ids = []
        if main_sub.get("PickUpCharacterCardId"):
            pickup_ids.append(main_sub.get("PickUpCharacterCardId"))
        pickup_ids.extend(main_sub.get("PickUpCharacterCardIds", []) or [])

        story_meta = event_story.get(event_id, {})
        story_sections = []
        for row in sorted(stories.get(event_id, []), key=lambda r: r.get("EventStorySectionNo", 0)):
            story_sections.append({
                "No": row.get("EventStorySectionNo"),
                "Group": row.get("EventStorySectionGroup"),
                "Type": row.get("EventStorySectionType"),
                "Title": row.get("TitleName", ""),
                "ReleaseDateTime": row.get("ReleaseDateTime", ""),
                "RequiredValue": row.get("KeyTargetValue", 0),
                "KeyCost": row.get("KeyEventStoryKeyUnlockCost", 0),
                "HasVoice": row.get("HasVoiceFile", False),
                "PopupText": clean_text(row.get("PopupText", "")),
                "ReadRewards": rewards.direct(row.get("ReadDirectRewardGroupId")),
            })

        rule_sections = []
        for row in sorted(rules.get(event_id, []), key=lambda r: r.get("SlideNo", 0)):
            rule_sections.append({
                "SlideNo": row.get("SlideNo"),
                "Type": row.get("RuleWindowSlideType"),
                "ReleaseDateTime": row.get("ReleaseDateTime", ""),
                "DescriptionFileName": row.get("DescriptionFileName", ""),
                "Description": clean_text(row.get("Description", "")),
            })

        shift_sections = []
        for row in sorted(shifts.get(event_id, []), key=lambda r: r.get("ShiftId", 0)):
            shift_sections.append({
                "ShiftId": row.get("ShiftId"),
                "CharacterId": row.get("CharacterId"),
                "CharacterName": character_map.get(row.get("CharacterId"), ""),
                "StartTime": row.get("StartTime", ""),
                "EndTime": row.get("EndTime", ""),
                "SpecialTimes": sorted(special_by_shift.get(row.get("ShiftId"), []), key=lambda r: r.get("SpecialTimeSequenceNo", 0)),
            })
        for row in sorted(ojt_shifts.get(event_id, []), key=lambda r: r.get("OjtShiftId", 0)):
            shift_sections.append({
                "ShiftId": row.get("OjtShiftId"),
                "CharacterId": row.get("CharacterId"),
                "CharacterName": character_map.get(row.get("CharacterId"), ""),
                "StartTime": row.get("StartTime", ""),
                "EndTime": row.get("EndTime", ""),
                "TrainingStartScript": row.get("TrainingPuzzleStartStoryFileName", ""),
                "TrainingClearScript": row.get("TrainingPuzzleClearStoryFileName", ""),
            })

        reward_summary = {
            "Ranking": [
                {
                    "RankingType": r.get("RankingType"),
                    "StartRank": r.get("StartRank"),
                    "WappenCount": r.get("RewardWappenCount"),
                    "Rewards": rewards.present(r.get("PresentId")),
                }
                for r in sorted(ranking_rewards.get(event_id, []), key=lambda r: (r.get("RankingType", 0), r.get("StartRank", 0)))
            ],
            "Sales": [
                {
                    "ShiftId": r.get("ShiftId"),
                    "KeySales": r.get("KeySales"),
                    "IsPickUp": r.get("IsPickUp"),
                    "Rewards": rewards.direct(r.get("DirectRewardGroupId")),
                }
                for r in sorted(sales_rewards.get(event_id, []), key=lambda r: (r.get("ShiftId", 0), r.get("KeySales", 0)))
            ],
            "Accumulate": [
                {
                    "KeyCount": r.get("KeyCount"),
                    "IsPickUp": r.get("IsPickUp"),
                    "Rewards": rewards.direct(r.get("DirectRewardGroupId")),
                }
                for r in sorted(accumulate_rewards.get(event_id, []), key=lambda r: r.get("KeyCount", 0))
            ],
        }
        event_format = EVENT_FORMAT_MAP.get(event.get("EventFormat"), {"Key": "unknown", "Label": "未知活动类型"})
        subtype_label = SOURCE_SUBTYPE_MAP.get(subtype, SOURCE_SUBTYPE_MAP["unknown"])
        classification = classify_event(event, event_a.get(event_id))

        archive.append({
            "EventId": event_id,
            "Title": event.get("EventTitle", ""),
            "EventFormat": event.get("EventFormat"),
            "EventFormatName": event_format["Key"],
            "EventFormatLabel": event_format["Label"],
            "ActivityType": classification["DisplayLabel"],
            "Classification": classification,
            "Availability": dict(date_window(
                event.get('OpenStartTime') or event.get('StartTime'), event.get('EndTime'),
                as_of=as_of, active=event.get('IsActive', True)),
                StartField='OpenStartTime' if event.get('OpenStartTime') else 'StartTime',
                EndField='EndTime'),
            "SourceSubtype": subtype,
            "SourceSubtypeLabel": subtype_label,
            "Subtype": subtype,
            "SubtypeLabel": subtype_label,
            "StartTime": event.get("StartTime", ""),
            "OpenStartTime": event.get("OpenStartTime", ""),
            "SecondHalfStartTime": event.get("SecondHalfStartTime", ""),
            "RankingEndTime": event.get("EventRankingEndTime", ""),
            "HighScoreRankingEndTime": event.get("HighScoreRankingEndTime", ""),
            "AnnouncementStartTime": event.get("AnnouncementStartTime", ""),
            "EndTime": event.get("EndTime", ""),
            "NoticeId": event.get("NoticeId", 0),
            "Assets": {
                "Logo": event.get("LogoFileName", ""),
                "PromotionalPopup": event.get("PromotionalPopupFileName", ""),
                "PuzzleTerminalBackground": event.get("PuzzleTerminalBackgroundFileName", ""),
                "PuzzleTerminalCharacters": event.get("PuzzleTerminalCharacterFileNames", []),
                "StoryBanner": story_meta.get("BannerFileName", ""),
                "StoryButton": story_meta.get("StoryButtonFileName", ""),
            },
            "Colors": {
                "Main": event.get("MainColor", ""),
                "Accent": event.get("AccentColor", ""),
                "Spectrum": event.get("SpectrumColor", ""),
            },
            "Characters": [{"Id": cid, "Name": character_map.get(cid, f"Unknown({cid})")} for cid in character_ids],
            "Cards": [{"Id": cid, "Name": card_map.get(cid, f"Unknown({cid})")} for cid in card_ids],
            "PickUpCards": [{"Id": cid, "Name": card_map.get(cid, f"Unknown({cid})")} for cid in pickup_ids],
            "EventItems": _names([main_sub.get("EventItemId"), main_sub.get("TravelCoinItemId"), main_sub.get("EventStaminaItemId"), main_sub.get("PrizeMedalItemId"), main_sub.get("ChartStampItemId")], item_map),
            "ExchangeId": main_sub.get("ExchangeId", 0),
            "StorySynopsis": clean_text(story_meta.get("Synopsis", "")),
            "StoryUnlockStartTime": story_meta.get("UnlockStartTime", ""),
            "StorySections": story_sections,
            "Rules": rule_sections,
            "Shifts": shift_sections,
            "Recipes": [
                {
                    "RecipeId": r.get("RecipeId"),
                    "RecipeName": r.get("RecipeName", ""),
                    "RecommendCharacter": character_map.get(r.get("RecommendCharacterId"), ""),
                    "IngredientIds": r.get("IngredientIds", []),
                    "Memo": clean_text(r.get("RecipeMemo", "")),
                    "FileName": r.get("RecipeFileName", ""),
                }
                for r in sorted(recipes.get(event_id, []), key=lambda r: r.get("RecipeId", 0))
            ],
            "PuzzleIngredientStageCount": len(ingredient_stages.get(event_id, [])),
            "Rewards": reward_summary,
            "Missions": event_missions.get(event_id, []),
        })

    out = json_path("Event_Archive.json")
    save_json({"Events": archive, "TimeAssessment": {
        'AsOf': as_of.isoformat(), 'Timezone': 'UTC', 'Boundary': '[start,end)',
        'AccountUnlock': 'not_assessed', 'DateFiltering': 'none',
    }}, out)
    save_json({"Issues": rewards.issues}, audit_path("event_reward_resolution.json"))
    save_json(mission_audit, audit_path("event_mission_relations.json"))
    if mission_audit['Issues']:
        record_warning(f"活动任务有 {len(mission_audit['Issues'])} 条关系或奖励缺项，详见 event_mission_relations.json")
    if rewards.issues:
        record_warning(f"活动奖励有 {len(rewards.issues)} 条解析或名称缺项，详见 event_reward_resolution.json")
    print(f"[+] extracted {len(archive)} events -> {out}")
    return archive


def _join_named(items):
    return " / ".join(item.get("Name", "") for item in items if item.get("Name"))


def _reward_text(rewards):
    parts = []
    for reward in rewards or []:
        name = reward.get("RewardName") or reward.get("RewardType")
        count = reward.get("RewardCount")
        label = reward.get("RewardTypeLabel") or reward.get("RewardType")
        parts.append(f"{name} x{count} [{label}]")
    return " / ".join(parts)


def export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    data = load_json(json_path("Event_Archive.json"))
    events = data.get("Events", [])
    save_json({"Archive": data, "Mappings": {
        "EventFormat": EVENT_FORMAT_MAP,
        "SourceSubtype": SOURCE_SUBTYPE_MAP,
        "RewardType": REWARD_TYPE_MAP,
    }}, audit_path("event_archive_audit.json"))
    out = xlsx_path("event_archive.xlsx")

    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    overview.append([
        "活动ID", "活动名", "活动类型",
        "开始", "开放", "后半开放", "排名结束", "结束", "关联角色", "关联卡牌",
        "PickUp卡牌", "兑换所ID", "剧情章节数", "规则页数", "活动材料关卡数", "Logo", "弹窗图", "日期状态",
    ])
    for e in events:
        overview.append([
            e.get("EventId"), e.get("Title"), e.get("ActivityType"), e.get("StartTime"),
            e.get("OpenStartTime"), e.get("SecondHalfStartTime"),
            e.get("RankingEndTime"), e.get("EndTime"), _join_named(e.get("Characters", [])),
            _join_named(e.get("Cards", [])), _join_named(e.get("PickUpCards", [])),
            e.get("ExchangeId"), len(e.get("StorySections", [])), len(e.get("Rules", [])),
            e.get("PuzzleIngredientStageCount"), e.get("Assets", {}).get("Logo"),
            e.get("Assets", {}).get("PromotionalPopup"),
            STATUS_LABELS.get(e.get('Availability', {}).get('Status'), '时间状态未知'),
        ])

    story = wb.create_sheet("Story")
    story.append(["活动ID", "活动名", "章节序号", "章节组", "章节标题", "开放时间", "解锁需求值", "钥匙消耗", "是否有语音", "弹窗文本", "阅读奖励"])
    for e in events:
        for s in e.get("StorySections", []):
            story.append([
                e.get("EventId"), e.get("Title"), s.get("No"), s.get("Group"),
                s.get("Title"), s.get("ReleaseDateTime"), s.get("RequiredValue"), s.get("KeyCost"),
                s.get("HasVoice"), s.get("PopupText"), _reward_text(s.get("ReadRewards")),
            ])

    rules = wb.create_sheet("Rules")
    rules.append(["活动ID", "活动名", "规则页序号", "开放时间", "说明图文件", "说明文本"])
    for e in events:
        for r in e.get("Rules", []):
            rules.append([
                e.get("EventId"), e.get("Title"), r.get("SlideNo"),
                r.get("ReleaseDateTime"), r.get("DescriptionFileName"), r.get("Description"),
            ])

    shifts = wb.create_sheet("Shifts")
    shifts.append(["活动ID", "活动名", "轮班/OJT ID", "角色", "开始", "结束", "训练开始脚本", "训练完成脚本"])
    for e in events:
        for s in e.get("Shifts", []):
            shifts.append([
                e.get("EventId"), e.get("Title"), s.get("ShiftId"), s.get("CharacterName"),
                s.get("StartTime"), s.get("EndTime"), s.get("TrainingStartScript", ""), s.get("TrainingClearScript", ""),
            ])

    rewards = wb.create_sheet("RewardSummary")
    rewards.append(["活动ID", "活动名", "奖励类别", "条件/档位", "是否重点奖励", "奖励内容"])
    for e in events:
        for r in e.get("Rewards", {}).get("Ranking", []):
            rewards.append([e.get("EventId"), e.get("Title"), "Ranking", f"type {r.get('RankingType')} rank {r.get('StartRank')}", "", _reward_text(r.get("Rewards"))])
        for r in e.get("Rewards", {}).get("Sales", []):
            rewards.append([e.get("EventId"), e.get("Title"), "Sales", f"shift {r.get('ShiftId')} sales {r.get('KeySales')}", r.get("IsPickUp"), _reward_text(r.get("Rewards"))])
        for r in e.get("Rewards", {}).get("Accumulate", []):
            rewards.append([e.get("EventId"), e.get("Title"), "Accumulate", r.get("KeyCount"), r.get("IsPickUp"), _reward_text(r.get("Rewards"))])

    missions = wb.create_sheet('EventMissions')
    missions.append(['活动ID', '活动名', '任务ID', '阶段', '目标值', '任务内容',
                     '开始', '结束', '隐藏', '仅计数', '奖励内容', '日期状态'])
    for event in events:
        for mission in event.get('Missions', []):
            raw = mission['RawMission']
            for seq in mission['Sequences']:
                missions.append([event['EventId'], event['Title'], mission['MissionId'],
                                 seq['SequenceNo'], seq['Border'], seq['Description'],
                                 raw.get('StartTime', ''), raw.get('EndTime', ''),
                                 seq['IsHidden'], seq['OnlyAccounting'], _reward_text(seq['Rewards']),
                                 STATUS_LABELS.get(mission.get('Availability', {}).get('Status'), '时间状态未知')])

    timing = wb.create_sheet('TimeAssessment')
    timing.append(['项目', '说明'])
    timing.append(['核对时刻', data.get('TimeAssessment', {}).get('AsOf', '')])
    timing.append(['时区', 'UTC；原始起止时间保留各自偏移'])
    timing.append(['区间边界', '包含开始，不包含结束'])
    timing.append(['含义', '日期区间内不代表账号已解锁；事件使用开放时间至结束时间，并非排名期'])
    timing.append(['收录', '不按日期筛除历史、未来或占位记录'])

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in col[:80]) + 2, 60)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and ("<br>" in cell.value or "\n" in cell.value):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = save_workbook_safely(wb, out)
    print(f"  [xlsx] {out}")
    return out


def run(session=None, *, as_of=None):
    extract(session=session, as_of=as_of)
    export()
