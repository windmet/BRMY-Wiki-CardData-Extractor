"""Annual character birthday archive, separate from cycle-selected dialogue."""
from collections import defaultdict, Counter

from openpyxl import Workbook

from ..core.availability import assessment_time, date_window, STATUS_LABELS
from ..core.exporter import audit_path, xlsx_path, save_workbook_safely
from ..core.scanner import save_json
from ..core.rewards import RewardResolver
from ..core.output import record_warning


def extract(session, *, as_of=None):
    tables = session.tables
    now = assessment_time(as_of)
    resolver = RewardResolver(tables)
    issues = []
    def index(name):
        groups = defaultdict(list)
        for row in tables.rows(name, active_only=True):
            groups[(row.get('CharacterId'), row.get('Year'))].append(row)
        return groups
    names = {'CampaignPage': 'mst_character_birthday_campaign_page',
             'LoginPage': 'mst_character_birthday_login_bonus_page',
             'MiniGame': 'mst_character_birthday_mini_game',
             'Texts': 'mst_character_birthday_mini_game_text',
             'Login': 'mst_character_birthday_login_bonus_sequence'}
    groups = {key: index(name) for key, name in names.items()}
    masters = index('mst_character_birthday')
    characters = tables.by_id('mst_character', 'CharacterId', required=False)
    costumes = tables.group_by('mst_costume_model', 'CostumeModelId', required=False, active_only=False)
    records = []
    for key, rows in sorted(masters.items()):
        if len(rows) != 1:
            issues.append({'Status': 'ambiguous_birthday', 'Key': key, 'Rows': rows})
            continue
        raw = rows[0]; cid, year = key
        entry = {'CharacterId': cid, 'Year': year,
                 'CharacterName': characters.get(cid, {}).get('CharacterNameJpn', f'Character({cid})'),
                 'Raw': raw, 'LoginRewards': [], 'TapRewards': [], 'Costumes': [],
                 'Availability': date_window(raw.get('StartTime'), raw.get('EndTime'), as_of=now)}
        for field in ('CampaignPage', 'LoginPage', 'MiniGame', 'Texts'):
            entry[field] = groups[field].get(key, [])
            if field != 'Texts' and len(entry[field]) > 1:
                issues.append({'Status': 'ambiguous_birthday_page', 'Key': key, 'Table': names[field], 'Rows': entry[field]})
        for phase in ('FirstDay', 'AfterBirthday'):
            base_id = raw.get('CampaignCostumeModelId' + phase)
            page_ids = [p.get('CostumeModelId' + phase) for p in entry['CampaignPage']]
            ids = {i for i in [base_id, *page_ids] if i}
            if not ids:
                continue
            target = next(iter(ids)) if len(ids) == 1 else None
            matches = costumes.get(target, []) if target else []
            status = 'resolved' if len(matches) == 1 else 'missing_costume'
            if len(ids) > 1:
                status = 'conflicting_costume'
            elif len(matches) > 1:
                status = 'ambiguous_costume'
            name = matches[0].get('CostumeModelName') if len(matches) == 1 else ''
            if status == 'resolved' and not name:
                status = 'costume_name_unavailable'
            relation = {'Phase': phase, 'CostumeModelId': target, 'Name': name or f'Costume({target})',
                        'Status': status, 'BaseId': base_id, 'PageIds': page_ids, 'Records': matches}
            entry['Costumes'].append(relation)
            if status != 'resolved':
                issues.append({'Status': status, 'Key': key, 'Relation': relation})
        counts = Counter(r.get('Sequence') for r in groups['Login'].get(key, []))
        for row in sorted(groups['Login'].get(key, []), key=lambda r: r.get('Sequence', 0)):
            if counts[row.get('Sequence')] != 1:
                issues.append({'Status': 'ambiguous_login_sequence', 'Key': key, 'Raw': row})
                continue
            entry['LoginRewards'].append({'Raw': row, 'Rewards': resolver.present(row.get('PresentId'))})
        for field in ('TapRewardNo1', 'TapRewardNo2', 'TapRewardNo3', 'TapRewardSecretPin'):
            entry['TapRewards'].append({'Field': field, 'GroupId': raw.get(field),
                                        'Rewards': resolver.direct(raw.get(field))})
        records.append(entry)
    for field, grouped in groups.items():
        for key, rows in grouped.items():
            if key not in masters:
                issues.append({'Status': 'orphan_birthday_relation', 'Table': names[field], 'Key': key, 'Rows': rows})
    return {'Birthdays': records, 'AsOf': now.isoformat(), 'Timezone': 'UTC',
            'Issues': issues, 'RewardIssues': resolver.issues,
            'RawTables': {n: tables.rows(n) for n in tables.names if n.startswith('mst_character_birthday')}}


def export(data):
    wb = Workbook(); overview = wb.active; overview.title = 'AnnualOverview'
    overview.append(['角色ID', '角色名', '年份', '开始', '结束', '日期状态', '登录奖励阶段数', '小游戏配置数'])
    login = wb.create_sheet('LoginRewards'); login.append(['角色ID', '角色名', '年份', '阶段', '奖励'])
    tap = wb.create_sheet('TapRewards'); tap.append(['角色ID', '角色名', '年份', '奖励位置', '奖励'])
    texts = wb.create_sheet('Texts'); texts.append(['角色ID', '角色名', '年份', '台词编号', '目标值', '台词', '播放时长原值'])
    costumes = wb.create_sheet('DisplayCostumes'); costumes.append(['角色ID', '角色名', '年份', '展示阶段', '服装ID', '服装名', '关联状态'])
    def reward_text(rows):
        return ' / '.join(f"{r['RewardName']} x{r['RewardCount']}" for r in rows)
    for entry in data['Birthdays']:
        prefix = [entry['CharacterId'], entry['CharacterName'], entry['Year']]; raw = entry['Raw']
        overview.append(prefix + [raw.get('StartTime'), raw.get('EndTime'), STATUS_LABELS[entry['Availability']['Status']], len(entry['LoginRewards']), len(entry['MiniGame'])])
        for row in entry['LoginRewards']:
            login.append(prefix + [row['Raw'].get('Sequence'), reward_text(row['Rewards'])])
        for row in entry['TapRewards']:
            if row['GroupId']:
                tap.append(prefix + [row['Field'], reward_text(row['Rewards'])])
        for row in entry['Texts']:
            texts.append(prefix + [row.get(k) for k in ('CharacterBirthdayTextNo', 'KeyTargetValue', 'Text', 'TextPlayTime')])
        for row in entry.get('Costumes', []):
            costumes.append(prefix + ['首日' if row['Phase'] == 'FirstDay' else '生日后', row['CostumeModelId'], row['Name'], row['Status']])
    notes = wb.create_sheet('Notes'); notes.append(['核对时刻 UTC', data['AsOf']])
    notes.append(['收录范围', '全部活跃角色年份记录，未按日期过滤；不替代原生日周期台词表'])
    notes.append(['日期状态', '日期区间内不代表账号已解锁'])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or '')) for c in column[:80]) + 2))
    return save_workbook_safely(wb, xlsx_path('birthday_archive.xlsx'))


def run(session=None, *, as_of=None):
    data = extract(session, as_of=as_of)
    save_json(data, audit_path('birthday_archive.json'))
    if data['Issues'] or data['RewardIssues']:
        record_warning(f"年度生日有 {len(data['Issues'])} 条关系异常、{len(data['RewardIssues'])} 条奖励缺项，详见 birthday_archive.json")
    return export(data)
