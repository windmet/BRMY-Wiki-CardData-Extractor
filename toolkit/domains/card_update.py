"""Merge a newly extracted card sheet with Wiki-maintained workbook columns."""
from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.exporter import (
    json_path,
    save_workbook_safely,
    write_workbook,
    xlsx_path,
)
from ..core.console import safe_print
from ..core.scanner import load_json, save_json
from . import cards
from .card_export import CARD_MANUAL_HEADERS, build_card_sheet


KEY_HEADER = "卡牌编号"


@dataclass
class OldCardTable:
    workbook: object
    sheet: object
    headers: list
    records: dict
    cells: dict
    row_numbers: dict


def _card_key(value):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"无效卡牌编号: {value!r}")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str) and value.strip().isdigit():
        return str(int(value.strip()))
    raise ValueError(f"无效卡牌编号: {value!r}")


def _read_old_table(path):
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["cards_data"] if "cards_data" in workbook.sheetnames else workbook.active
    raw_headers = [cell.value for cell in sheet[1]]
    while raw_headers and raw_headers[-1] in (None, ""):
        raw_headers.pop()
    if not raw_headers or KEY_HEADER not in raw_headers:
        raise ValueError(f"旧工作簿缺少主键列: {KEY_HEADER}")
    if any(header in (None, "") for header in raw_headers):
        raise ValueError("旧工作簿表头中存在空列")
    if len(set(raw_headers)) != len(raw_headers):
        raise ValueError("旧工作簿存在重复列名")

    key_column = raw_headers.index(KEY_HEADER) + 1
    records = {}
    cells = {}
    row_numbers = {}
    for row_number in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=row_number, column=index + 1).value for index in range(len(raw_headers))]
        if all(value in (None, "") for value in values):
            continue
        key = _card_key(sheet.cell(row=row_number, column=key_column).value)
        if key in records:
            raise ValueError(f"旧工作簿存在重复卡牌编号: {key}")
        records[key] = dict(zip(raw_headers, values))
        cells[key] = {
            header: sheet.cell(row=row_number, column=index + 1)
            for index, header in enumerate(raw_headers)
        }
        row_numbers[key] = row_number
    return OldCardTable(workbook, sheet, raw_headers, records, cells, row_numbers)


def _new_table(headers, rows):
    if KEY_HEADER not in headers:
        raise ValueError(f"新卡牌表缺少主键列: {KEY_HEADER}")
    if len(set(headers)) != len(headers):
        raise ValueError("新卡牌表存在重复列名")
    records = {}
    order = []
    for row in rows:
        record = dict(zip(headers, row))
        key = _card_key(record.get(KEY_HEADER))
        if key in records:
            raise ValueError(f"新卡牌表存在重复卡牌编号: {key}")
        records[key] = record
        order.append(key)
    return records, order


def _same_value(left, right):
    left = None if left == "" else left
    right = None if right == "" else right
    return left == right


def _copy_cell(source, target):
    target.value = "" if source.value is None and source.data_type == "inlineStr" else source.value
    _copy_cell_format(source, target)


def _copy_cell_format(source, target):
    if source.has_style:
        # StyleArray contains indexes into the owning workbook's tables.
        # Assign components so openpyxl registers them in the new workbook.
        for component in ("font", "fill", "border", "alignment", "protection"):
            setattr(target, component, copy(getattr(source, component)))
        target.quotePrefix = source.quotePrefix
        target.pivotButton = source.pivotButton
    target.number_format = source.number_format
    if source.comment:
        target.comment = copy(source.comment)
    if source.hyperlink:
        target.hyperlink = copy(source.hyperlink)


def _style_header(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="3F4B5B")
    cell.alignment = Alignment(vertical="center")


def merge_card_workbook(
    old_path,
    new_headers,
    new_rows,
    *,
    updated_path,
    changes_path,
    audit_path,
):
    old = _read_old_table(old_path)
    new_records, new_order = _new_table(new_headers, new_rows)
    old_keys = set(old.records)
    new_keys = set(new_records)

    manual_headers = [header for header in CARD_MANUAL_HEADERS if header in new_headers]
    custom_headers = [header for header in old.headers if header not in new_headers]
    preserved_headers = [
        header for header in manual_headers if header in old.headers
    ] + custom_headers
    final_headers = list(new_headers) + custom_headers

    added = [key for key in new_order if key not in old_keys]
    removed = [key for key in old.records if key not in new_keys]
    comparable_headers = [
        header
        for header in new_headers
        if header != KEY_HEADER and header not in manual_headers
    ]
    modified = []
    modified_cards = set()
    for key in new_order:
        if key not in old.records:
            continue
        old_record = old.records[key]
        new_record = new_records[key]
        for header in comparable_headers:
            old_value = old_record.get(header)
            new_value = new_record.get(header)
            if not _same_value(old_value, new_value):
                modified.append([int(key), header, old_value, new_value])
                modified_cards.add(key)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "cards_data"
    sheet.append(final_headers)
    for index, header in enumerate(final_headers, start=1):
        target = sheet.cell(row=1, column=index)
        if header in old.headers:
            source = old.sheet.cell(row=1, column=old.headers.index(header) + 1)
            _copy_cell(source, target)
        else:
            _style_header(target)
    sheet.freeze_panes = old.sheet.freeze_panes or "A2"

    preserved_cells = 0
    formula_cells = []
    final_rows = {}
    for output_row, key in enumerate(new_order, start=2):
        record = dict(new_records[key])
        record[KEY_HEADER] = int(key)
        for header in custom_headers:
            record[header] = ""
        sheet.append([record.get(header, "") for header in final_headers])
        if key in old.records:
            for header in old.headers:
                if header not in final_headers:
                    continue
                source = old.cells[key][header]
                target = sheet.cell(row=output_row, column=final_headers.index(header) + 1)
                _copy_cell_format(source, target)
            for header in preserved_headers:
                source = old.cells[key][header]
                target = sheet.cell(row=output_row, column=final_headers.index(header) + 1)
                _copy_cell(source, target)
                record[header] = source.value
                if source.value not in (None, ""):
                    preserved_cells += 1
                    if source.data_type == "f":
                        formula_cells.append({"CardId": int(key), "Column": header})
        else:
            sheet.cell(
                row=output_row,
                column=final_headers.index(KEY_HEADER) + 1,
            ).number_format = "0"
        final_rows[key] = [record.get(header, "") for header in final_headers]

    for output_index, header in enumerate(final_headers, start=1):
        if header not in old.headers:
            continue
        old_letter = get_column_letter(old.headers.index(header) + 1)
        new_letter = get_column_letter(output_index)
        width = old.sheet.column_dimensions[old_letter].width
        if width:
            sheet.column_dimensions[new_letter].width = width
    sheet.auto_filter.ref = sheet.dimensions
    actual_updated_path = save_workbook_safely(workbook, updated_path)
    print(f"  [xlsx] {actual_updated_path}")

    summary_rows = [
        ["旧表卡牌", len(old.records)],
        ["新表卡牌", len(new_records)],
        ["新增卡牌", len(added)],
        ["自动字段有修改的卡牌", len(modified_cards)],
        ["修改单元格", len(modified)],
        ["移除卡牌", len(removed)],
        ["保留人工/自定义非空单元格", preserved_cells],
        ["保留公式单元格", len(formula_cells)],
    ]
    actual_changes_path = write_workbook(changes_path, [
        {"title": "Summary", "headers": ["指标", "数量"], "rows": summary_rows},
        {
            "title": "Added",
            "headers": final_headers,
            "rows": [final_rows[key] for key in added],
        },
        {
            "title": "Modified",
            "headers": [KEY_HEADER, "字段", "旧值", "新值"],
            "rows": modified,
        },
        {
            "title": "Removed",
            "headers": old.headers,
            "rows": [
                [old.records[key].get(header, "") for header in old.headers]
                for key in removed
            ],
        },
    ])

    audit = {
        "Status": "PASS",
        "SourceWorkbook": str(Path(old_path).resolve()),
        "KeyHeader": KEY_HEADER,
        "ManualColumns": manual_headers,
        "ManualColumnsFound": [header for header in manual_headers if header in old.headers],
        "ManualColumnsMissing": [header for header in manual_headers if header not in old.headers],
        "CustomColumns": custom_headers,
        "Counts": {
            "OldCards": len(old.records),
            "NewCards": len(new_records),
            "AddedCards": len(added),
            "ModifiedCards": len(modified_cards),
            "ModifiedCells": len(modified),
            "RemovedCards": len(removed),
            "PreservedNonEmptyCells": preserved_cells,
            "PreservedFormulaCells": len(formula_cells),
        },
        "AddedCardIds": [int(key) for key in added],
        "ModifiedCardIds": [int(key) for key in sorted(modified_cards, key=int)],
        "RemovedCardIds": [int(key) for key in removed],
        "PreservedFormulaCells": formula_cells,
        "Outputs": {
            "UpdatedWorkbook": str(Path(actual_updated_path).resolve()),
            "ChangesWorkbook": str(Path(actual_changes_path).resolve()),
        },
    }
    save_json(audit, audit_path)
    old.workbook.close()
    safe_print(
        f"[+] 卡牌增量更新：新增 {len(added)}，"
        f"修改 {len(modified_cards)}，移除 {len(removed)}"
    )
    return audit


def run(old_workbook, audio_dir=None, session=None):
    old_workbook = str(Path(old_workbook).resolve())
    if not Path(old_workbook).is_file():
        raise FileNotFoundError(f"旧卡牌工作簿不存在: {old_workbook}")
    cards.extract(audio_dir=audio_dir, session=session)
    data = load_json(json_path("All_Cards_Database.json"))
    headers, rows = build_card_sheet(data)
    return merge_card_workbook(
        old_workbook,
        headers,
        rows,
        updated_path=xlsx_path("cards_data_updated.xlsx"),
        changes_path=xlsx_path("cards_data_changes.xlsx"),
        audit_path=json_path("cards_update_audit.json"),
    )
