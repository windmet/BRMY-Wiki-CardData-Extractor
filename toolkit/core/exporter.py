"""通用导出：XLSX + 输出目录管理。"""
import os
import re
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    OPENPYXL = True
except ImportError:
    OPENPYXL = False

JSON_DIR = 'json_output'
XLSX_DIR = 'xlsx_output'


def ensure_dirs():
    os.makedirs(JSON_DIR, exist_ok=True)
    os.makedirs(XLSX_DIR, exist_ok=True)


def json_path(filename):
    ensure_dirs()
    return os.path.join(JSON_DIR, filename)


def xlsx_path(filename):
    ensure_dirs()
    return os.path.join(XLSX_DIR, filename)


def _coerce_excel_value(value, column_type, column_index):
    if column_type is None or value is None:
        return value
    if column_type == "text":
        return str(value)
    if column_type == "integer":
        if isinstance(value, bool):
            raise ValueError(f"column {column_index + 1} expects integer, got bool")
        if isinstance(value, int):
            return value
        if isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value):
            return int(value)
        raise ValueError(
            f"column {column_index + 1} expects integer, got {value!r}"
        )
    raise ValueError(f"unsupported Excel column type: {column_type!r}")


def write_xlsx(
    rows,
    path,
    headers=None,
    sheet_title="Sheet1",
    col_widths=None,
    wrap_cols=None,
    column_types=None,
):
    """Write one sheet while preserving values unless a zero-based column type is declared."""
    if not OPENPYXL:
        print("  [!] openpyxl 未安装，跳过 xlsx 生成")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    if headers:
        ws.append(headers)

    column_types = dict(column_types or {})
    for column_index, column_type in column_types.items():
        if not isinstance(column_index, int) or column_index < 0:
            raise ValueError("Excel column type indexes must be non-negative integers")
        if column_type not in {"text", "integer"}:
            raise ValueError(f"unsupported Excel column type: {column_type!r}")
    for row in rows:
        cleaned = [
            _coerce_excel_value(item, column_types.get(index), index)
            for index, item in enumerate(row)
        ]
        ws.append(cleaned)

    first_data_row = 2 if headers else 1
    for column_index, column_type in column_types.items():
        number_format = "@" if column_type == "text" else "0"
        for row_index in range(first_data_row, ws.max_row + 1):
            ws.cell(row=row_index, column=column_index + 1).number_format = number_format

    if col_widths:
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

    if wrap_cols:
        for r_idx in range(2, ws.max_row + 1):
            for c_idx in wrap_cols:
                cell = ws.cell(row=r_idx, column=c_idx + 1)
                if cell.value and isinstance(cell.value, str) and '<br>' in str(cell.value):
                    cell.alignment = Alignment(wrap_text=True, vertical='center')

    wb.save(path)
    print(f"  [xlsx] {path}")


def write_workbook(path, sheets):
    """Write a workbook from sheet dictionaries with consistent table ergonomics."""
    if not OPENPYXL:
        print("  [!] openpyxl 未安装，跳过 xlsx 生成")
        return False

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["title"])
        headers = sheet.get("headers", [])
        if headers:
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="3F4B5B")
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"

        for row in sheet.get("rows", []):
            ws.append(list(row))

        if headers and ws.max_row >= 1:
            ws.auto_filter.ref = ws.dimensions
        for column, width in sheet.get("col_widths", {}).items():
            ws.column_dimensions[column].width = width
        for column_index in sheet.get("wrap_cols", []):
            for row_index in range(2, ws.max_row + 1):
                ws.cell(row=row_index, column=column_index).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    actual_path = path
    try:
        wb.save(actual_path)
    except PermissionError:
        stem, extension = os.path.splitext(path)
        actual_path = f"{stem}_new{extension}"
        if os.path.exists(actual_path):
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            actual_path = f"{stem}_new_{timestamp}{extension}"
        wb.save(actual_path)
        print(f"  [!] 原 XLSX 正被占用，已改存: {actual_path}")
    print(f"  [xlsx] {actual_path}")
    return actual_path
